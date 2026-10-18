"""
_build_excel.py builds greek_banking_model.xlsx.

Every figure is read with pandas from
02_Banking_Sector_Dashboard/data/processed/kpis_final.csv (one row per bank and year).
The report and PDF page of each figure are listed in sources.csv in the same folder.

A few results cannot be recomputed from that file (EBA 2023 stress test, ECB funding,
2022 underlying profit). They are typed in VERIFIED below, next to the 2026 NII base case
and the CAMELS-style composite. Every result that can be recomputed is checked against
VERIFIED before the workbook is written, so the build stops if the data and the verified
values disagree.

Run from the deliverables folder:   python3 _build_excel.py
"""
import re
import sys
from math import ceil
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parent.parent
PROCESSED = ROOT / "02_Banking_Sector_Dashboard" / "data" / "processed"
KPIS_CSV = PROCESSED / "kpis_final.csv"
SOURCES_CSV = PROCESSED / "sources.csv"
ECB_RATES_CSV = ROOT / "02_Banking_Sector_Dashboard" / "data" / "manual" / "ecb_deposit_rate.csv"
OUT = Path(__file__).resolve().parent / "greek_banking_model.xlsx"

BANKS = ["Eurobank", "Alpha Bank", "Piraeus Bank", "NBG"]
YEARS = [2022, 2023, 2024]

ENTITIES = {
    "Alpha Bank": "Alpha Services and Holdings S.A.",
    "Eurobank": "Eurobank Ergasias Services and Holdings S.A.",
    "NBG": "National Bank of Greece S.A.",
    "Piraeus Bank": "Piraeus Financial Holdings S.A.",
}

BANK_COLORS = {
    "Eurobank": "0067B1",
    "Alpha Bank": "E2231A",
    "Piraeus Bank": "F7A600",
    "NBG": "003087",
}

# Stated assumptions
COE = 10.0   # cost of equity, %
G = 2.0      # long-run growth, %
STRESS = {"cor_bp": 200, "loans": -0.15, "nim_bp": -50, "tax": 0.29}
ASSET_GROWTH_PCT = 2

# ---------------------------------------------------------------------------
# Verified results (typed values only for what kpis_final.csv cannot give)
# ---------------------------------------------------------------------------
VERIFIED = {
    "sector_nii": {2022: 5376, 2023: 8099, 2024: 8596},
    "sector_nii_growth_pct": 59.9,
    "sector_net_profit": {2022: 3785, 2023: 3652, 2024: 4326},
    "walk_2023": {"delta": 2723, "rate": 2891},
    "walk_2024": {"delta": 497, "rate": 270, "volume": 225},
    "eurobank_volume_2024": 265,
    "sector_underlying_2022": 2577,          # not recomputable from the CSV
    "reported_change_2023_pct": -3.5,
    "underlying_change_2023_pct": 42,        # "about 42%"
    # net profit (EUR m), CET1 end-2024, CET1 end of stress year, extra cost of risk (bp)
    "stress": {
        "NBG": (371, 18.3, 19.3, 420),
        "Eurobank": (402, 16.8, 17.6, 367),
        "Piraeus Bank": (217, 14.5, 15.1, 353),
        "Alpha Bank": (-118, 16.3, 15.9, 209),
    },
    # official EBA 2023 stress test, fully loaded CET1 (%): Dec 2022, adverse end-2025
    "eba_2023": {
        "Alpha Bank": (11.9, 8.9),
        "Eurobank": (14.4, 12.2),
        "NBG": (15.9, 14.5),
        "Piraeus Bank": (11.5, 9.1),
    },
    "nii_2026_base": {"Eurobank": 2430, "Alpha Bank": 1434, "Piraeus Bank": 1828, "NBG": 1902},
    "sector_nii_2026_base": 7593,
    "sector_nii_2026_vs_2024_pct": -11.7,
    "dtc_total_bn": {2022: 13.5, 2024: 12.0},
    "dtc_to_cet1_2024": {"Eurobank": 36, "Piraeus Bank": 63},
    "ecb_funding_bn": {2022: 35.1, 2023: 14.3, 2024: 2.6},   # all Alpha in 2024
    "camels_2024": {"Eurobank": 1.8, "NBG": 1.8, "Piraeus Bank": 2.2, "Alpha Bank": 3.0},
}

# ---------------------------------------------------------------------------
# Data and analysis (the same block is used in _build_pptx.py)
# ---------------------------------------------------------------------------
kpis = pd.read_csv(KPIS_CSV).sort_values(["bank", "year"]).reset_index(drop=True)
sources = pd.read_csv(SOURCES_CSV)
assert len(kpis) == len(BANKS) * len(YEARS), "kpis_final.csv must have one row per bank and year"
assert set(kpis.bank) == set(BANKS) and set(kpis.year) == set(YEARS)
kpis["ppop"] = kpis["operating_income"] + kpis["operating_expenses"]   # profit before impairments

k24 = kpis[kpis.year == 2024].set_index("bank").loc[BANKS]
sector = kpis.groupby("year")[["nii", "operating_income", "net_profit", "loans", "deposits",
                               "total_assets", "equity", "dtc"]].sum()


def rate_volume_walk(df):
    """Change in NII split into rate (margin), volume and mix, per bank and year."""
    rows = []
    for bank in BANKS:
        b = df[df.bank == bank].set_index("year")
        nim = b["nii"] / b["avg_total_assets"]          # unrounded NIM, so the walk adds up
        for yr in (2023, 2024):
            p = yr - 1
            d_nim = nim[yr] - nim[p]
            d_a = b.loc[yr, "avg_total_assets"] - b.loc[p, "avg_total_assets"]
            rows.append({
                "bank": bank, "year": yr,
                "delta": b.loc[yr, "nii"] - b.loc[p, "nii"],
                "rate": d_nim * b.loc[p, "avg_total_assets"],
                "volume": nim[p] * d_a,
                "mix": d_nim * d_a,
            })
    w = pd.DataFrame(rows)
    assert ((w.rate + w.volume + w.mix) - w.delta).abs().max() < 0.5
    return w


def stress_test(df24):
    """One-year stress on FY2024: +200bp cost of risk, -15% loans, -50bp NIM."""
    rows = {}
    for bank in BANKS:
        r = df24.loc[bank]
        nim = r["nii"] / r["avg_total_assets"]
        cor = -r["impairment"] / r["avg_loans"]
        cet1_capital = r["cet1"] / 100 * r["rwa"]
        loans_adv = r["loans"] * (1 + STRESS["loans"])
        assets_adv = r["avg_total_assets"] + r["loans"] * STRESS["loans"]
        nii_adv = (nim + STRESS["nim_bp"] / 10_000) * assets_adv
        imp_adv = (cor + STRESS["cor_bp"] / 10_000) * loans_adv
        pbt_adv = r["pbt"] + (nii_adv - r["nii"]) - (imp_adv + r["impairment"])
        np_adv = r["net_profit"] + (pbt_adv - r["pbt"]) * (1 - STRESS["tax"])
        rows[bank] = {
            "net_profit": r["net_profit"], "net_profit_stress": np_adv,
            "cet1_start": r["cet1"], "cet1_end": (cet1_capital + np_adv) / r["rwa"] * 100,
            "absorb_bp": r["pbt"] / r["loans"] * 10_000,
        }
    return pd.DataFrame(rows).T.loc[BANKS]


def justified_ptbv(rote):
    return (rote - G) / (COE - G)


# Cut-offs of the CAMELS-style score (1 = best, 5 = worst): the project's own, not a supervisory method
CAMELS_CUTS = {
    "cet1": ([11.0, 13.0, 15.0, 18.0], "higher"),
    "npe_ratio": ([2.0, 3.0, 5.0, 8.0], "lower"),
    "cost_to_income": ([30.0, 37.0, 45.0, 55.0], "lower"),
    "roe": ([5.0, 8.0, 12.0, 15.0], "higher"),
    "loan_to_deposit": ([60.0, 65.0, 75.0, 85.0], "lower"),
}


def camels_score(value, cuts, direction):
    t = sorted(cuts)
    if direction == "higher":
        return 1 + sum(value < x for x in reversed(t))
    return 1 + sum(value > x for x in t)


def camels_composite(df24):
    cols = {m: [camels_score(df24.loc[b, m], *CAMELS_CUTS[m]) for b in BANKS] for m in CAMELS_CUTS}
    return pd.Series(pd.DataFrame(cols, index=BANKS).mean(axis=1).round(1), index=BANKS)


def nii_2026_base_case():
    """NII 2026 = NIM x average assets. NIM = 2024 NIM + slope x (2026 rate - 2024 rate), the slope being
    the bank's 2022-2024 regression of NIM on the yearly average ECB deposit rate. Assets grow 2% a year
    from end-2024. Used only to check the typed base case values."""
    ecb = pd.read_csv(ECB_RATES_CSV).set_index("year")["dfr_avg"]
    out = {}
    for bank in BANKS:
        b = kpis[kpis.bank == bank].sort_values("year")
        slope = np.polyfit([ecb[y] for y in YEARS], b["nim"].tolist(), 1)[0]
        nim26 = b["nim"].iloc[-1] + slope * (2.00 - ecb[2024])
        assets26 = b["total_assets"].iloc[-1] * 1.02 * 1.01
        out[bank] = nim26 / 100 * assets26
    return pd.Series(out)


walk = rate_volume_walk(kpis)
stress = stress_test(k24)
walk_sector = walk.groupby("year")[["delta", "rate", "volume", "mix"]].sum()
nii_growth = (sector.loc[2024, "nii"] / sector.loc[2022, "nii"] - 1) * 100
np_2023_reported = (sector.loc[2023, "net_profit"] / sector.loc[2022, "net_profit"] - 1) * 100
np_2023_underlying = (sector.loc[2023, "net_profit"] / VERIFIED["sector_underlying_2022"] - 1) * 100


def verify_results():
    """Stop the build if a recomputed result differs from the verified value."""
    V = VERIFIED
    for y in YEARS:
        assert round(sector.loc[y, "nii"]) == V["sector_nii"][y], f"sector NII {y}"
        assert round(sector.loc[y, "net_profit"]) == V["sector_net_profit"][y], f"sector profit {y}"
    assert round(nii_growth, 1) == V["sector_nii_growth_pct"]
    assert round(walk_sector.loc[2023, "delta"]) == V["walk_2023"]["delta"]
    assert round(walk_sector.loc[2023, "rate"]) == V["walk_2023"]["rate"]
    assert round(walk_sector.loc[2024, "delta"]) == V["walk_2024"]["delta"]
    assert round(walk_sector.loc[2024, "rate"]) == V["walk_2024"]["rate"]
    assert round(walk_sector.loc[2024, "volume"]) == V["walk_2024"]["volume"]
    eb = walk[(walk.bank == "Eurobank") & (walk.year == 2024)].iloc[0]
    assert round(eb["volume"]) == V["eurobank_volume_2024"]
    assert round(np_2023_reported, 1) == V["reported_change_2023_pct"]
    assert round(np_2023_underlying) == V["underlying_change_2023_pct"]
    for bank, (np_s, c0, c1, bp) in V["stress"].items():
        r = stress.loc[bank]
        assert round(r["net_profit_stress"]) == np_s, f"stress profit {bank}"
        assert round(r["cet1_start"], 1) == c0, f"stress CET1 start {bank}"
        assert round(r["cet1_end"], 1) == c1, f"stress CET1 end {bank}"
        assert round(r["absorb_bp"]) == bp, f"stress absorb {bank}"
    for y, v in V["dtc_total_bn"].items():
        assert round(sector.loc[y, "dtc"] / 1000, 1) == v, f"DTC total {y}"
    for bank, v in V["dtc_to_cet1_2024"].items():
        assert round(k24.loc[bank, "dtc_to_cet1"]) == v, f"DTC share {bank}"
    assert sum(V["nii_2026_base"].values()) - V["sector_nii_2026_base"] in (0, 1)   # rounding
    assert round((V["sector_nii_2026_base"] / sector.loc[2024, "nii"] - 1) * 100, 1) \
        == V["sector_nii_2026_vs_2024_pct"]
    base_check = nii_2026_base_case()
    for bank, v in V["nii_2026_base"].items():
        assert abs(base_check[bank] - v) <= 0.6, f"NII 2026 base {bank}"
    assert abs(base_check.sum() - V["sector_nii_2026_base"]) <= 0.6
    comp = camels_composite(k24)
    for bank, v in V["camels_2024"].items():
        assert abs(comp[bank] - v) < 1e-9, f"CAMELS composite {bank}"


verify_results()

# ---------------------------------------------------------------------------
# Style
# ---------------------------------------------------------------------------
NAVY = "0A1628"
DARK_BLUE = "1B3A6B"
MID_BLUE = "2D5A9E"
LIGHTEST = "EEF4FC"
WHITE = "FFFFFF"
GREEN = "1A7340"
LIGHT_GREEN = "D6F0E0"
RED = "991B1B"
GREY_BG = "F1F5F9"
GREY_TXT = "64748B"
INK = "0F172A"
BORDER_COL = "CBD5E1"


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def border():
    s = Side(style="thin", color=BORDER_COL)
    return Border(left=s, right=s, top=s, bottom=s)


def put(ws, row, col, value, *, bold=False, color=INK, size=10, italic=False, bg=None,
        h="left", v="center", wrap=False, fmt=None, bord=False, indent=0):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name="Calibri", bold=bold, italic=italic, size=size, color=color)
    c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)
    if bg:
        c.fill = fill(bg)
    if fmt:
        c.number_format = fmt
    if bord:
        c.border = border()
    return c


def section_bar(ws, row, title, last_col, bg=DARK_BLUE, height=18):
    ws.row_dimensions[row].height = height
    put(ws, row, 1, title, bold=True, color=WHITE, bg=bg)
    for col in range(2, last_col + 1):
        ws.cell(row=row, column=col).fill = fill(bg)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)


def note_row(ws, row, text, last_col, chars_per_line=150):
    lines = max(1, ceil(len(text) / chars_per_line))
    ws.row_dimensions[row].height = 13.5 * lines + 3
    put(ws, row, 1, text, italic=True, size=9, color=GREY_TXT, wrap=True, v="top")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)


def est_height(text, chars_per_line, line_h=14.0):
    return max(1, ceil(len(text) / chars_per_line)) * line_h + 4


# Number formats. Ratios are stored as fractions and shown as percentages.
F_INT = "#,##0"
F_P1 = "0.0%"
F_P2 = "0.00%"
F_PP1 = '▲ 0.0" pp";▼ 0.0" pp";0.0" pp"'
F_PP2 = '▲ 0.00" pp";▼ 0.00" pp";0.00" pp"'
F_CH = "▲ 0.0%;▼ 0.0%;0.0%"
F_X = '0.00"x"'
F_BP = '#,##0" bp"'


def ratio(v):
    return None if v is None or pd.isna(v) else v / 100


# ---------------------------------------------------------------------------
# Source text from sources.csv (report and page of each statement line or ratio)
# ---------------------------------------------------------------------------
SOURCE_METRIC = {
    "nii": "Net interest income",
    "fees": "Net banking fee and commission income",
    "operating_income": "Operating income",
    "operating_expenses": "Operating expenses",
    "ppop": "Profit from operations before impairments",
    "impairment": "Impairment losses on loans",
    "pbt": "Profit before tax",
    "net_profit": "Net profit attributable to shareholders",
    "loans": "Loans and advances to customers",
    "deposits": "Deposits from customers",
    "total_assets": "Total assets",
    "equity": "Equity attributable to shareholders",
    "at1": "Additional Tier 1 instruments",
    "intangibles": "Goodwill and intangible assets",
    "cet1": "cet1_ratio",
    "total_capital_ratio": "total_capital_ratio",
    "rwa": "rwa",
    "dtc": "dtc",
    "npe_ratio": "npe_ratio",
}
DERIVED_SOURCE = {
    "tangible_equity": "Derived: equity less AT1 less goodwill and intangibles",
    "dtc_to_cet1": "Derived: DTC / (CET1 ratio × RWA)",
}
RATIO_SOURCE = "Calculated from the figures above (definitions on the Assumptions sheet)"


def source_text(bank, key):
    if key in DERIVED_SOURCE:
        return DERIVED_SOURCE[key]
    if key not in SOURCE_METRIC:
        return RATIO_SOURCE
    rows = sources[(sources.bank == bank) & (sources.metric == SOURCE_METRIC[key])
                   & (sources.year.isin(YEARS))].sort_values("year")
    groups = {}
    for _, r in rows.iterrows():
        groups.setdefault(r["source"], []).append(int(r["year"]))
    parts = []
    for src, yrs in groups.items():
        m = re.match(r"^.+_(\d{4})\.pdf p\.(\d+)$", src)
        assert m, f"unexpected source format: {src}"
        parts.append(f"FY{m.group(1)} report p.{m.group(2)} ({', '.join(str(y) for y in yrs)})")
    return "; ".join(parts)


wb = Workbook()
wb.properties.creator = "Spyros Papastergiou"
wb.properties.lastModifiedBy = "Spyros Papastergiou"
wb.properties.title = "Greek Banking Sector Analysis 2022–2024"
wb.properties.description = "Excel model built from kpis_final.csv. Not investment advice."

SHEET_ORDER = [
    ("Cover", "This sheet: contents, basis and contact"),
    ("Sector Summary", "Cross-bank KPIs 2022–2024, with sector totals"),
    ("Eurobank", "Income statement, balance sheet, returns and capital, with report pages"),
    ("Alpha Bank", "Income statement, balance sheet, returns and capital, with report pages"),
    ("Piraeus Bank", "Income statement, balance sheet, returns and capital, with report pages"),
    ("NBG", "Income statement, balance sheet, returns and capital, with report pages"),
    ("Analysis", "NII rate and volume walk, 2022 one-offs, stress test, NII scenarios, valuation, DTCs, ECB funding"),
    ("Assumptions", "Entities, definitions and the parameters of the stress test, scenarios and valuation"),
]

# ===========================================================================
# SHEET 1: COVER
# ===========================================================================
ws = wb.active
ws.title = "Cover"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 4
ws.column_dimensions["B"].width = 20
ws.column_dimensions["C"].width = 100
ws.column_dimensions["D"].width = 4
for row in range(1, 30):
    for col in range(1, 5):
        ws.cell(row=row, column=col).fill = fill(NAVY)

ws.row_dimensions[2].height = 50
ws.row_dimensions[3].height = 30
put(ws, 2, 2, "GREEK BANKING SECTOR ANALYSIS", bold=True, size=22, color=WHITE, bg=NAVY)
put(ws, 3, 2, "Eurobank · Alpha Bank · Piraeus Bank · NBG  |  FY2022–2024", size=13,
    color="94A3B8", bg=NAVY)
put(ws, 9, 2, "TABLE OF CONTENTS", bold=True, size=12, color="60A5FA", bg=NAVY)
for i, (sheet, desc) in enumerate(SHEET_ORDER, 11):
    put(ws, i, 2, sheet, bold=True, color=WHITE, bg=NAVY)
    put(ws, i, 3, desc, color="94A3B8", bg=NAVY)

info = [
    ("Analyst:", "Spyros Papastergiou"),
    ("Contact:", "papastergiousp@gmail.com"),
    ("GitHub:", "github.com/papastergiousp-maker/greek-banking-sector-analysis"),
    ("Dashboard:", "phytai.com/dashboard"),
    ("Entities:", "The four companies listed on the Athens Exchange in 2022–2024, consolidated group figures: "
                  + ", ".join(ENTITIES[b] for b in ["Alpha Bank", "Eurobank", "NBG", "Piraeus Bank"])),
    ("Basis:", "Latest restated figures: 2024 and 2023 from the FY2024 annual reports (2023 restated), "
               "2022 from the FY2023 annual reports. Every number is traced to a PDF page (sources.csv)."),
    ("Disclaimer:", "For portfolio and educational purposes only. Not investment advice."),
]
r = 20
for label, text in info:
    put(ws, r, 2, label, bold=True, color="60A5FA", bg=NAVY, v="top")
    put(ws, r, 3, text, color=WHITE, bg=NAVY, wrap=True, v="top", italic=(label == "Disclaimer:"))
    ws.row_dimensions[r].height = est_height(text, 105)
    r += 1

# ===========================================================================
# SHEET 2: SECTOR SUMMARY
# ===========================================================================
SUMMARY = [
    # label, kpis column, number format, best direction in 2024, additive across banks
    ("INCOME AND PROFITABILITY", None, None, None, None),
    ("NII (€m)", "nii", F_INT, "high", True),
    ("Operating income (€m)", "operating_income", F_INT, "high", True),
    ("Net profit (€m)", "net_profit", F_INT, "high", True),
    ("ROE, average equity (%)", "roe", F_P1, "high", False),
    ("RoTE, average tangible equity (%)", "rote", F_P1, "high", False),
    ("ROA, average assets (%)", "roa", F_P2, "high", False),
    ("NIM, average assets (%)", "nim", F_P2, "high", False),
    ("EFFICIENCY AND RISK", None, None, None, None),
    ("Cost-to-income (%)", "cost_to_income", F_P1, "low", False),
    ("Cost of risk (%)", "cost_of_risk", F_P2, "low", False),
    ("BALANCE SHEET", None, None, None, None),
    ("Loans (€m)", "loans", F_INT, None, True),
    ("Deposits (€m)", "deposits", F_INT, None, True),
    ("Total assets (€m)", "total_assets", F_INT, None, True),
    ("Equity attributable (€m)", "equity", F_INT, None, True),
    ("Loan-to-deposit (%)", "loan_to_deposit", F_P1, None, False),
    ("CAPITAL AND ASSET QUALITY", None, None, None, None),
    ("CET1 ratio (%)", "cet1", F_P1, "high", False),
    ("Total capital ratio (%)", "total_capital_ratio", F_P1, "high", False),
    ("NPE ratio (%)", "npe_ratio", F_P1, "low", False),
    ("DTC (€m)", "dtc", F_INT, "low", True),
    ("DTC share of CET1 (%)", "dtc_to_cet1", F_P1, "low", False),
]


def make_sector_summary():
    ws = wb.create_sheet("Sector Summary")
    ws.sheet_view.showGridLines = False
    last_col = 17   # A + 12 bank-year columns + 3 sector columns + change column
    ws.column_dimensions["A"].width = 34
    for col in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 11.5
    ws.column_dimensions[get_column_letter(last_col)].width = 14

    ws.row_dimensions[1].height = 30
    put(ws, 1, 1, "SECTOR SUMMARY: KEY PERFORMANCE INDICATORS (FY2022–2024)", bold=True, size=13,
        color=WHITE, bg=DARK_BLUE)
    for col in range(2, last_col + 1):
        ws.cell(row=1, column=col).fill = fill(DARK_BLUE)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)

    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 18
    put(ws, 2, 1, "Metric", bold=True, color=WHITE, bg=DARK_BLUE)
    ws.merge_cells("A2:A3")

    col = 2
    for bank in BANKS:
        put(ws, 2, col, bank, bold=True, color=WHITE, bg=BANK_COLORS[bank], h="center", bord=True)
        ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + 2)
        for j, yr in enumerate(YEARS):
            put(ws, 3, col + j, yr, bold=True, size=9, color=WHITE, bg=MID_BLUE, h="center", bord=True)
        col += 3
    put(ws, 2, 14, "Four banks (sum)", bold=True, color=WHITE, bg=DARK_BLUE, h="center", bord=True)
    ws.merge_cells("N2:P2")
    for j, yr in enumerate(YEARS):
        put(ws, 3, 14 + j, yr, bold=True, size=9, color=WHITE, bg=MID_BLUE, h="center", bord=True)
    put(ws, 2, 17, "Change 2022 to 2024", bold=True, size=9, color=WHITE, bg=DARK_BLUE, h="center",
        wrap=True, bord=True)
    ws.merge_cells("Q2:Q3")

    row = 4
    for label, key, fmt, direction, additive in SUMMARY:
        ws.row_dimensions[row].height = 16
        if key is None:
            section_bar(ws, row, label, last_col, height=16)
            row += 1
            continue
        band = GREY_BG if row % 2 == 0 else WHITE
        put(ws, row, 1, label, bg=band, bord=True, indent=1)

        values_2024 = {b: k24.loc[b, key] for b in BANKS}
        best = None
        if direction:
            vals = list(values_2024.values())
            best = max(vals) if direction == "high" else min(vals)

        col = 2
        for bank in BANKS:
            for yr in YEARS:
                raw = kpis[(kpis.bank == bank) & (kpis.year == yr)][key].iloc[0]
                shown = raw / 100 if "%" in fmt else raw
                bg, fg, bold = (LIGHTEST if yr == 2024 else band), INK, False
                if yr == 2024 and best is not None and abs(raw - best) < 1e-9:
                    bg, fg, bold = LIGHT_GREEN, GREEN, True
                put(ws, row, col, shown, bg=bg, color=fg, bold=bold, h="center", fmt=fmt, bord=True)
                col += 1
        for j, yr in enumerate(YEARS):
            if additive:
                put(ws, row, 14 + j, sector.loc[yr, key], bg=LIGHTEST, h="center", fmt=fmt, bord=True,
                    bold=True)
            else:
                put(ws, row, 14 + j, None, bg=GREY_BG, bord=True)
        if additive:
            chg = sector.loc[2024, key] / sector.loc[2022, key] - 1
            put(ws, row, 17, chg, bg=LIGHTEST, h="center", fmt=F_CH, bord=True, bold=True)
        else:
            put(ws, row, 17, None, bg=GREY_BG, bord=True)
        row += 1

    row += 1
    note_row(ws, row, "Green = best of the four banks in 2024 (highest, or lowest for cost-to-income, cost of risk, "
                      "NPE ratio, DTC and DTC share of CET1). Size items and the loan-to-deposit ratio are not ranked. "
                      "Sector columns show the sum of the four banks for amounts only.", last_col, 190)
    note_row(ws, row + 1, "Eurobank 2024 includes Hellenic Bank, consolidated from 30 June 2024. Source: kpis_final.csv; "
                          "report and page of each figure on the bank sheets and in sources.csv.", last_col, 190)
    ws.freeze_panes = "B4"


make_sector_summary()

# ===========================================================================
# SHEETS 3 TO 6: ONE PER BANK
# ===========================================================================
# label, kpis column, kind (amount, cost, ratio), number format, good direction (+1, -1, 0)
IS_ROWS = [
    ("Net interest income", "nii", "amount", F_INT, 1),
    ("Net fee and commission income", "fees", "amount", F_INT, 1),
    ("Operating income", "operating_income", "amount", F_INT, 1),
    ("Operating expenses", "operating_expenses", "cost", F_INT, -1),
    ("Profit before impairments (PPOP)", "ppop", "amount", F_INT, 1),
    ("Impairment losses on loans", "impairment", "cost", F_INT, -1),
    ("Profit before tax", "pbt", "amount", F_INT, 1),
    ("Net profit attributable to shareholders", "net_profit", "amount", F_INT, 1),
]
BS_ROWS = [
    ("Loans and advances to customers", "loans", "amount", F_INT, 0),
    ("Customer deposits", "deposits", "amount", F_INT, 0),
    ("Total assets", "total_assets", "amount", F_INT, 0),
    ("Equity attributable to shareholders", "equity", "amount", F_INT, 0),
    ("   of which AT1 instruments", "at1", "amount", F_INT, 0),
    ("Goodwill and intangible assets", "intangibles", "amount", F_INT, 0),
    ("Tangible equity", "tangible_equity", "amount", F_INT, 0),
]
RET_ROWS = [
    ("ROE (average equity)", "roe", "ratio", F_P1, 1),
    ("RoTE (average tangible equity)", "rote", "ratio", F_P1, 1),
    ("ROA (average assets)", "roa", "ratio", F_P2, 1),
    ("NIM (average assets)", "nim", "ratio", F_P2, 1),
    ("Cost-to-income", "cost_to_income", "ratio", F_P1, -1),
    ("Cost of risk (average loans)", "cost_of_risk", "ratio", F_P2, -1),
    ("Fee share of operating income", "fee_share", "ratio", F_P1, 0),
    ("Loan-to-deposit", "loan_to_deposit", "ratio", F_P1, 0),
]
CAP_ROWS = [
    ("CET1 ratio", "cet1", "ratio", F_P1, 1),
    ("Total capital ratio", "total_capital_ratio", "ratio", F_P1, 1),
    ("Risk-weighted assets", "rwa", "amount", F_INT, 0),
    ("DTC included in capital", "dtc", "amount", F_INT, -1),
    ("DTC share of CET1 capital", "dtc_to_cet1", "ratio", F_P1, -1),
    ("NPE ratio", "npe_ratio", "ratio", F_P1, -1),
]


def change_value(kind, prev, cur):
    """Change between two years: percentage points for ratios, percent for amounts.
    Cost lines (stored as negatives) are compared on their size. None if the base is zero."""
    if kind == "ratio":
        return cur - prev
    base = abs(prev)
    if base == 0:
        return None
    return ((abs(cur) - abs(prev)) if kind == "cost" else (cur - prev)) / base


def make_bank_sheet(bank):
    bcolor = BANK_COLORS[bank]
    ws = wb.create_sheet(bank)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 40
    for col in range(2, 4 + 1):
        ws.column_dimensions[get_column_letter(col)].width = 12
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 78
    last_col = 7

    ws.row_dimensions[1].height = 30
    put(ws, 1, 1, f"{bank.upper()}: FINANCIAL SUMMARY (FY2022–2024)", bold=True, size=13, color=WHITE, bg=bcolor)
    for col in range(2, last_col + 1):
        ws.cell(row=1, column=col).fill = fill(bcolor)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)

    ws.row_dimensions[2].height = 18
    put(ws, 2, 1, f"{ENTITIES[bank]}, consolidated group. €m unless stated.", italic=True, size=9, color=GREY_TXT)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)

    ws.row_dimensions[3].height = 30
    heads = ["", 2022, 2023, 2024, "Change 2023 vs 2022", "Change 2024 vs 2023", "Report and page"]
    for j, h in enumerate(heads, 1):
        put(ws, 3, j, h, bold=True, size=10 if j < 5 else 9, color=WHITE, bg=MID_BLUE if j > 1 else DARK_BLUE,
            h="center" if j not in (1, 7) else "left", wrap=True, bord=True)

    bk = kpis[kpis.bank == bank].set_index("year")
    row = 4

    def block(title, spec, row):
        section_bar(ws, row, title, last_col, height=18)
        row += 1
        for label, key, kind, fmt, good in spec:
            ws.row_dimensions[row].height = 16
            band = GREY_BG if row % 2 == 0 else WHITE
            is_total = key in ("ppop", "net_profit", "tangible_equity")
            put(ws, row, 1, label, bold=is_total, bg=band, bord=True, indent=1)
            vals = {}
            for j, yr in enumerate(YEARS, 2):
                raw = float(bk.loc[yr, key])
                vals[yr] = raw
                shown = raw / 100 if kind == "ratio" else raw
                put(ws, row, j, shown, bold=is_total, bg=LIGHTEST if yr == 2024 else band, h="center",
                    fmt=fmt, bord=True)
            for col, (p, c) in zip((5, 6), ((2022, 2023), (2023, 2024))):
                val = change_value(kind, vals[p], vals[c])
                cfmt = F_CH if kind != "ratio" else (F_PP2 if fmt == F_P2 else F_PP1)
                color = GREY_TXT
                if good and val:
                    color = GREEN if val * good > 0 else RED
                if val is None:
                    put(ws, row, col, "n/a", size=9, color=GREY_TXT, bg=band, h="center", bord=True)
                else:
                    put(ws, row, col, val, size=9, color=color, bg=band, h="center", fmt=cfmt, bord=True)
            put(ws, row, 7, source_text(bank, key), size=8, color=GREY_TXT, bg=band, bord=True)
            row += 1
        return row

    row = block("INCOME STATEMENT (€m)", IS_ROWS, row)
    row += 1
    row = block("BALANCE SHEET (€m, year end)", BS_ROWS, row)
    row += 1
    row = block("RETURNS AND EFFICIENCY (%)", RET_ROWS, row)
    row += 1
    row = block("CAPITAL AND ASSET QUALITY", CAP_ROWS, row)
    row += 1
    extra = ""
    if bank == "Eurobank":
        extra = " Eurobank 2024 includes Hellenic Bank, consolidated from 30 June 2024."
    note_row(ws, row, "Change columns: percentage change for amounts (cost lines on their size), percentage-point "
                      "change for ratios. Green = better, red = worse, grey = not ranked. Averages are (opening + "
                      "closing) / 2." + extra, last_col, 190)
    ws.freeze_panes = "B4"


for bank in BANKS:
    make_bank_sheet(bank)

# ===========================================================================
# SHEET 7: ANALYSIS
# ===========================================================================


def add_table(ws, row, title, headers, rows, fmts, last_col=7, bold_last=False):
    section_bar(ws, row, title, last_col)
    row += 1
    lines = max(ceil(len(str(h)) / 15) for h in headers[1:])
    ws.row_dimensions[row].height = 14 * lines + 6
    for j, h in enumerate(headers, 1):
        put(ws, row, j, h, bold=True, size=9, color=WHITE, bg=MID_BLUE, h="left" if j == 1 else "center",
            wrap=True, bord=True)
    row += 1
    for i, vals in enumerate(rows):
        ws.row_dimensions[row].height = 16
        band = GREY_BG if i % 2 else WHITE
        is_last = bold_last and i == len(rows) - 1
        for j, v in enumerate(vals, 1):
            put(ws, row, j, v, bold=is_last, bg=band, h="left" if j == 1 else "center",
                fmt=fmts[j - 1], bord=True, indent=1 if j == 1 else 0)
        row += 1
    return row


def make_analysis():
    ws = wb.create_sheet("Analysis")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 46
    for col in range(2, 8):
        ws.column_dimensions[get_column_letter(col)].width = 17
    last_col = 7

    ws.row_dimensions[1].height = 30
    put(ws, 1, 1, "ANALYSIS: NII DRIVERS, EARNINGS QUALITY, STRESS TEST, SCENARIOS, VALUATION", bold=True,
        size=13, color=WHITE, bg=DARK_BLUE)
    for col in range(2, last_col + 1):
        ws.cell(row=1, column=col).fill = fill(DARK_BLUE)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    note_row(ws, 2, "Computed from kpis_final.csv unless a note names another source. €m unless stated. "
                    "Valuation is a way to read returns, not a price target. Not investment advice.", last_col)

    row = 4
    # 1. Sector totals
    rows = []
    for label, series, bn in (("Net interest income (€m)", sector["nii"], False),
                              ("Net profit attributable (€m)", sector["net_profit"], False),
                              ("DTC included in CET1 capital (€m)", sector["dtc"], False)):
        rows.append([label] + [series[y] for y in YEARS] + [series[2024] / series[2022] - 1])
    ecb = VERIFIED["ecb_funding_bn"]
    rows.append(["Due to the ECB and central banks (€bn, year end)"] + [ecb[y] for y in YEARS]
                + [ecb[2024] / ecb[2022] - 1])
    row = add_table(ws, row, "1. SECTOR TOTALS (SUM OF THE FOUR BANKS)",
                    ["Item", 2022, 2023, 2024, "Change 2022 to 2024"], rows,
                    [None, F_INT, F_INT, F_INT, F_CH])
    # ECB row uses one decimal (€bn)
    for col in (2, 3, 4):
        ws.cell(row=row - 1, column=col).number_format = "0.0"
    note_row(ws, row, "DTC = deferred tax credits: a tax claim on the Greek State that counts in CET1 capital. "
                      "ECB funding is the amount due to the ECB and central banks at year end (02_Banking_Sector_Dashboard/data/manual/ecb_funding.csv, "
                      "report and page in that file); all of the 2024 amount is Alpha's.", last_col)
    row += 2

    # 2. NII rate / volume walk
    for yr in (2023, 2024):
        rows = []
        for bank in BANKS:
            r = walk[(walk.bank == bank) & (walk.year == yr)].iloc[0]
            rows.append([bank, r["delta"], r["rate"], r["volume"], r["mix"]])
        s = walk_sector.loc[yr]
        rows.append(["Sector", s["delta"], s["rate"], s["volume"], s["mix"]])
        row = add_table(ws, row, f"2{'A' if yr == 2023 else 'B'}. NII RATE AND VOLUME WALK, {yr - 1} TO {yr} (€m)",
                        ["Bank", "Change in NII", "Rate (margin) effect", "Volume effect", "Mix"], rows,
                        [None, F_INT, F_INT, F_INT, F_INT], bold_last=True)
        row += 1
    note_row(ws, row - 1, "Rate effect = change in NIM × prior-year average assets. Volume effect = prior-year NIM × change "
                          "in average assets. Mix = change in NIM × change in average assets. The three add up to the change in "
                          "NII. Eurobank consolidated Hellenic Bank from 30 June 2024.", last_col)
    row += 1

    # 3. Earnings quality
    ws.row_breaks.append(Break(id=row - 1))
    rows = [
        ["Sector net profit 2022, reported (€m)", sector.loc[2022, "net_profit"]],
        ["Sector net profit 2022, underlying (€m)", VERIFIED["sector_underlying_2022"]],
        ["Sector net profit 2023, reported (€m)", sector.loc[2023, "net_profit"]],
        ["2023 vs 2022, reported", np_2023_reported / 100],
        ["2023 vs 2022, on underlying 2022 profit", np_2023_underlying / 100],
    ]
    row = add_table(ws, row, "3. EARNINGS QUALITY: 2022 ONE-OFFS", ["Item", "Value"], rows, [None, F_INT])
    ws.cell(row=row - 2, column=2).number_format = F_CH
    ws.cell(row=row - 1, column=2).number_format = F_CH
    note_row(ws, row, "2022 reported profit includes one-off items. The underlying figure removes them (notebook "
                      "03_analysis/05_earnings_quality.ipynb) and cannot be recomputed from kpis_final.csv. On "
                      "underlying figures 2023 profit rose about 42% instead of falling 3.5%.", last_col)
    row += 2

    # 4. Stress test and EBA 2023
    rows = [[b, stress.loc[b, "net_profit"], stress.loc[b, "net_profit_stress"], stress.loc[b, "cet1_start"] / 100,
             stress.loc[b, "cet1_end"] / 100, stress.loc[b, "absorb_bp"]] for b in BANKS]
    row = add_table(ws, row, "4. ONE-YEAR STRESS TEST ON FY2024 (SIMPLIFIED)",
                    ["Bank", "Net profit FY2024 (€m)", "Stressed net profit (€m)", "CET1 end-2024",
                     "CET1 end of stress year", "Extra cost of risk before a pre-tax loss"], rows,
                    [None, F_INT, F_INT, F_P1, F_P1, F_BP], last_col=7)
    note_row(ws, row, "Shocks: cost of risk up 200bp, loans down 15%, NIM down 50bp. Costs flat, 29% tax on the change in profit, "
                      "no dividends, RWAs unchanged. CET1 at the end of the stress year = (reported CET1 ratio × RWAs + "
                      "stressed net profit) / RWAs. A bank that stays profitable ends the year with more CET1 because no "
                      "dividend is paid. Extra cost of risk before a pre-tax loss = FY2024 profit before tax / loans. "
                      "Simplified, static, one year: directional only.", last_col)
    row += 2
    rows = [[b, VERIFIED["eba_2023"][b][0] / 100, VERIFIED["eba_2023"][b][1] / 100,
             VERIFIED["eba_2023"][b][1] - VERIFIED["eba_2023"][b][0]] for b in BANKS]
    row = add_table(ws, row, "OFFICIAL EBA 2023 STRESS TEST (FULLY LOADED CET1)",
                    ["Bank", "Dec 2022", "Adverse, end-2025", "Change"], rows,
                    [None, F_P1, F_P1, F_PP1])
    note_row(ws, row, "Three-year adverse scenario on the 31.12.2022 balance sheet, as reported by each bank "
                      "(02_Banking_Sector_Dashboard/data/manual/eba_stress_test_2023.csv). Not comparable with the simplified one-year test above.",
             last_col)
    row += 2

    # 5. NII scenarios
    ws.row_breaks.append(Break(id=row - 1))
    base = VERIFIED["nii_2026_base"]
    rows = [[b, k24.loc[b, "nii"], base[b], base[b] / k24.loc[b, "nii"] - 1] for b in BANKS]
    rows.append(["Sector", sector.loc[2024, "nii"], VERIFIED["sector_nii_2026_base"],
                 VERIFIED["sector_nii_2026_vs_2024_pct"] / 100])
    row = add_table(ws, row, "5. NII SCENARIO: BASE CASE 2026 (ROUGH GUIDE ONLY)",
                    ["Bank", "NII 2024, actual (€m)", "NII 2026, base case (€m)", "Change vs 2024"], rows,
                    [None, F_INT, F_INT, F_CH], bold_last=True)
    note_row(ws, row, "Not a forecast. NIM moves with the ECB yearly average deposit rate (each bank's 2022–2024 "
                      "sensitivity) and total assets grow 2% a year from end-2024. The base case assumes gradual ECB cuts. "
                      "Source: notebook 04_forecasting/01_forecast_nii_2025_2026.ipynb. Banks may not add to the sector "
                      "total because of rounding.", last_col)
    row += 2

    # 6. Valuation
    rows = [[b, k24.loc[b, "net_profit"], k24.loc[b, "avg_tangible_equity"], k24.loc[b, "rote"] / 100,
             COE / 100, G / 100, justified_ptbv(k24.loc[b, "rote"])] for b in BANKS]
    row = add_table(ws, row, "6. VALUATION: JUSTIFIED P/TBV (NOT INVESTMENT ADVICE)",
                    ["Bank", "Net profit 2024 (€m)", "Average tangible equity (€m)", "RoTE 2024", "CoE (assumed)",
                     "g (assumed)", "Justified P/TBV"], rows,
                    [None, F_INT, F_INT, F_P1, F_P1, F_P1, F_X])
    note_row(ws, row, "Justified P/TBV = (RoTE - g) / (CoE - g), with CoE 10% and g 2% as stated assumptions. No market "
                      "prices are used. A bank that earns exactly its cost of equity is worth 1.0x tangible book value "
                      "in this model. It is a way to read returns, not a price target. Not investment advice.", last_col)
    row += 2

    # 7. CAMELS-style composite
    composite = camels_composite(k24)
    rows = [[b, composite[b]] for b in BANKS]
    row = add_table(ws, row, "7. CAMELS-STYLE COMPOSITE, 2024 (1 = BEST)", ["Bank", "Composite"], rows, [None, "0.0"])
    note_row(ws, row, "Average of five scores (CET1 ratio, NPE ratio, cost-to-income, ROE, loan-to-deposit ratio) on "
                      "the project's own cut-offs. Not a supervisory rating. Source: notebook "
                      "03_analysis/02_camels_scorecard.ipynb.", last_col)


make_analysis()

# ===========================================================================
# SHEET 8: ASSUMPTIONS
# ===========================================================================


def make_assumptions():
    ws = wb.create_sheet("Assumptions")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 95
    last_col = 4

    ws.row_dimensions[1].height = 28
    put(ws, 1, 1, "ENTITIES, DEFINITIONS AND ASSUMPTIONS", bold=True, size=13, color=WHITE, bg=DARK_BLUE)
    for col in range(2, last_col + 1):
        ws.cell(row=1, column=col).fill = fill(DARK_BLUE)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)

    # "text" blocks: item and one text (merged across B to D). "kv" blocks: item, value, unit, note.
    blocks = [
        ("ENTITIES (CONSOLIDATED GROUP FIGURES)", "text", ["Short name", "Entity"], [
            ["Alpha Bank", ENTITIES["Alpha Bank"]],
            ["Eurobank", ENTITIES["Eurobank"] + " (includes Hellenic Bank from 30 June 2024)"],
            ["NBG", ENTITIES["NBG"]],
            ["Piraeus Bank", ENTITIES["Piraeus Bank"]],
        ]),
        ("DEFINITIONS", "text", ["Metric", "Definition"], [
            ["ROE", "Net profit attributable / average equity attributable"],
            ["RoTE", "Net profit / average (equity less AT1 less goodwill and intangibles). AT1 coupons are not deducted"],
            ["ROA", "Net profit / average total assets"],
            ["NIM", "Net interest income / average total assets"],
            ["Cost of risk", "Loan impairments / average loans"],
            ["Cost-to-income", "Operating expenses / operating income"],
            ["Loan-to-deposit", "Loans / customer deposits, year end"],
            ["CET1, total capital, NPE ratio", "As reported at each year end"],
            ["DTC", "Deferred tax credits included in regulatory capital: a tax claim on the Greek State counted in CET1"],
            ["DTC share of CET1", "DTC / (CET1 ratio × RWAs)"],
            ["Averages", "(Opening + closing) / 2"],
        ]),
        ("STRESS TEST (ONE YEAR, FY2024 BASE)", "kv", ["Shock", "Value", "Unit", "Note"], [
            ["Additional cost of risk", STRESS["cor_bp"], "bp", "On the loan book"],
            ["Loan volume", round(STRESS["loans"] * 100), "%", "Loans fall by 15%. Average total assets fall by the same euro amount"],
            ["NIM", STRESS["nim_bp"], "bp", "Margin compression"],
            ["Operating costs", "Flat", "", "No cost cutting"],
            ["Tax", round(STRESS["tax"] * 100), "%", "Applied to the change in profit before tax"],
            ["Dividends", "None", "", "No dividend in the stress year"],
            ["RWAs", "Unchanged", "", "Conservative: a smaller loan book would lower RWAs"],
        ]),
        ("NII SCENARIOS (ROUGH GUIDE ONLY)", "kv", ["Item", "Value", "Unit", "Note"], [
            ["Model", "NIM × average assets", "", "NIM and assets are projected separately and multiplied"],
            ["NIM driver", "ECB deposit rate", "", "Each bank's 2022–2024 sensitivity to the yearly average deposit rate. Three yearly points per bank, so directional only"],
            ["Total assets growth", ASSET_GROWTH_PCT, "% a year", "From the end-2024 balance sheet, the same in every scenario"],
            ["Base case", "Gradual ECB cuts", "", "Results on the Analysis sheet"],
        ]),
        ("VALUATION (NOT INVESTMENT ADVICE)", "kv", ["Parameter", "Value", "Unit", "Note"], [
            ["Cost of equity (CoE)", COE, "%", "Stated assumption"],
            ["Long-run growth (g)", G, "%", "Stated assumption"],
            ["Justified P/TBV", "(RoTE - g) / (CoE - g)", "", "No market prices are used. A way to read returns, not a price target"],
        ]),
        ("CAMELS-STYLE COMPOSITE", "text", ["Item", "Note"], [
            ["Dimensions", "C: CET1 ratio. A: NPE ratio. M: cost-to-income. E: ROE. L: loan-to-deposit ratio. S is left out"],
            ["Scale", "1 = best. Own cut-offs, not a supervisory rating"],
        ]),
        ("DATA CHECKS AND LIMITS", "text", ["Item", "Note"], [
            ["Automated tests", "tests/test_kpis.py recomputes the ratios from the statement lines and checks that assets equal liabilities plus equity"],
            ["Traceability", "sources.csv lists the report file and PDF page of every figure"],
            ["Operating income and expenses", "Each bank's own presentation, so cost-to-income ratios are close but not perfectly comparable"],
            ["CET1 ratios", "Include year-end profit net of the dividend accrual that each bank applies, so they are not perfectly comparable"],
            ["Data frequency", "Yearly only, so intra-year effects cannot be measured"],
        ]),
    ]

    row = 3
    for title, kind, headers, rows in blocks:
        section_bar(ws, row, title, last_col)
        row += 1
        ws.row_dimensions[row].height = 16
        if kind == "text":
            put(ws, row, 1, headers[0], bold=True, size=9, color=WHITE, bg=MID_BLUE, bord=True)
            put(ws, row, 2, headers[1], bold=True, size=9, color=WHITE, bg=MID_BLUE, bord=True)
            for col in (3, 4):
                ws.cell(row=row, column=col).fill = fill(MID_BLUE)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        else:
            for j, h in enumerate(headers, 1):
                put(ws, row, j, h, bold=True, size=9, color=WHITE, bg=MID_BLUE,
                    h="left" if j in (1, 4) else "center", bord=True)
        row += 1
        for i, vals in enumerate(rows):
            band = GREY_BG if i % 2 else WHITE
            text = str(vals[-1])
            ws.row_dimensions[row].height = est_height(text, 115, 13.5) if len(text) > 105 else 16
            if kind == "text":
                put(ws, row, 1, vals[0], bg=band, bord=True, indent=1)
                put(ws, row, 2, vals[1], bg=band, bord=True, wrap=True)
                for col in (3, 4):
                    ws.cell(row=row, column=col).fill = fill(band)
                ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            else:
                for j, v in enumerate(vals, 1):
                    put(ws, row, j, v if v != "" else None, bg=band, h="left" if j in (1, 4) else "center",
                        bord=True, wrap=(j == 4), indent=1 if j == 1 else 0)
            row += 1
        row += 1


make_assumptions()

assert wb.sheetnames == [n for n, _ in SHEET_ORDER], wb.sheetnames
for sheet in wb.worksheets:                      # print setup: landscape, one page wide
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1 if sheet.title == "Cover" else 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT)
print(f"Written: {OUT}  ({len(wb.sheetnames)} sheets: {', '.join(wb.sheetnames)})")
