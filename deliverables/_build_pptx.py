"""
_build_pptx.py builds greek_banking_deck.pptx (12 slides).

Every figure is read with pandas from
02_Banking_Sector_Dashboard/data/processed/kpis_final.csv (one row per bank and year).
The report and PDF page of each figure are listed in sources.csv in the same folder.

A few results cannot be recomputed from that file (EBA 2023 stress test, ECB funding,
2022 underlying profit). They are typed in VERIFIED below, next to the 2026 NII base case
and the CAMELS-style composite. Every result that can be recomputed is checked against
VERIFIED before the deck is written, so the build stops if the data and the verified
values disagree. The analysis block is the same as in _build_excel.py.

Run from the deliverables folder:   python3 _build_pptx.py
"""
import io
import sys
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from lxml import etree
from matplotlib.lines import Line2D
from matplotlib.patches import Patch
from matplotlib.ticker import FuncFormatter
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parent.parent
PROCESSED = ROOT / "02_Banking_Sector_Dashboard" / "data" / "processed"
KPIS_CSV = PROCESSED / "kpis_final.csv"
ECB_RATES_CSV = ROOT / "02_Banking_Sector_Dashboard" / "data" / "manual" / "ecb_deposit_rate.csv"
OUT = Path(__file__).resolve().parent / "greek_banking_deck.pptx"

BANKS = ["Eurobank", "Alpha Bank", "Piraeus Bank", "NBG"]
YEARS = [2022, 2023, 2024]
N_SLIDES = 12

ENTITIES = {
    "Alpha Bank": "Alpha Services and Holdings S.A.",
    "Eurobank": "Eurobank Ergasias Services and Holdings S.A.",
    "NBG": "National Bank of Greece S.A.",
    "Piraeus Bank": "Piraeus Financial Holdings S.A.",
}

CONTACT_EMAIL = "papastergiousp@gmail.com"
CONTACT_GITHUB = "github.com/papastergiousp-maker/greek-banking-sector-analysis"
CONTACT_DASHBOARD = "phytai.com/dashboard"
NOT_ADVICE = "Not investment advice."

# Stated assumptions
COE = 10.0   # cost of equity, %
G = 2.0      # long-run growth, %
STRESS = {"cor_bp": 200, "loans": -0.15, "nim_bp": -50, "tax": 0.29}

# ---------------------------------------------------------------------------
# Verified results (typed values only for what kpis_final.csv cannot give)
# ---------------------------------------------------------------------------
VERIFIED = {
    "sector_nii": {2022: 5376, 2023: 8099, 2024: 8596},
    "sector_nii_growth_pct": 59.9,
    "sector_net_profit": {2022: 3785, 2023: 3652, 2024: 4326},
    "walk_2023": {"delta": 2723, "rate": 2891},
    "walk_2024": {"delta": 497, "rate": 270, "volume": 225},
    "eurobank_volume_2024": 265,
    "sector_underlying_2022": 2577,          # not recomputable from the CSV
    "reported_change_2023_pct": -3.5,
    "underlying_change_2023_pct": 42,        # "about 42%"
    # net profit (EUR m), CET1 end-2024, CET1 end of stress year, extra cost of risk (bp)
    "stress": {
        "NBG": (371, 18.3, 19.3, 420),
        "Eurobank": (402, 16.8, 17.6, 367),
        "Piraeus Bank": (217, 14.5, 15.1, 353),
        "Alpha Bank": (-118, 16.3, 15.9, 209),
    },
    # official EBA 2023 stress test, fully loaded CET1 (%): Dec 2022, adverse end-2025
    "eba_2023": {
        "Alpha Bank": (11.9, 8.9),
        "Eurobank": (14.4, 12.2),
        "NBG": (15.9, 14.5),
        "Piraeus Bank": (11.5, 9.1),
    },
    "nii_2026_base": {"Eurobank": 2430, "Alpha Bank": 1434, "Piraeus Bank": 1828, "NBG": 1902},
    "sector_nii_2026_base": 7593,
    "sector_nii_2026_vs_2024_pct": -11.7,
    "dtc_total_bn": {2022: 13.5, 2024: 12.0},
    "dtc_to_cet1_2024": {"Eurobank": 36, "Piraeus Bank": 63},
    "ecb_funding_bn": {2022: 35.1, 2023: 14.3, 2024: 2.6},   # all Alpha in 2024
    "camels_2024": {"Eurobank": 1.8, "NBG": 1.8, "Piraeus Bank": 2.2, "Alpha Bank": 3.0},
}

# ---------------------------------------------------------------------------
# Data and analysis (the same block is used in _build_excel.py)
# ---------------------------------------------------------------------------
kpis = pd.read_csv(KPIS_CSV).sort_values(["bank", "year"]).reset_index(drop=True)
assert len(kpis) == len(BANKS) * len(YEARS), "kpis_final.csv must have one row per bank and year"
assert set(kpis.bank) == set(BANKS) and set(kpis.year) == set(YEARS)

k24 = kpis[kpis.year == 2024].set_index("bank").loc[BANKS]
k23 = kpis[kpis.year == 2023].set_index("bank").loc[BANKS]
k22 = kpis[kpis.year == 2022].set_index("bank").loc[BANKS]
sector = kpis.groupby("year")[["nii", "operating_income", "net_profit", "loans", "deposits",
                               "total_assets", "equity", "dtc"]].sum()


def rate_volume_walk(df):
    """Change in NII split into rate (margin), volume and mix, per bank and year."""
    rows = []
    for bank in BANKS:
        b = df[df.bank == bank].set_index("year")
        nim = b["nii"] / b["avg_total_assets"]          # unrounded NIM, so the walk adds up
        for yr in (2023, 2024):
            p = yr - 1
            d_nim = nim[yr] - nim[p]
            d_a = b.loc[yr, "avg_total_assets"] - b.loc[p, "avg_total_assets"]
            rows.append({
                "bank": bank, "year": yr,
                "delta": b.loc[yr, "nii"] - b.loc[p, "nii"],
                "rate": d_nim * b.loc[p, "avg_total_assets"],
                "volume": nim[p] * d_a,
                "mix": d_nim * d_a,
            })
    w = pd.DataFrame(rows)
    assert ((w.rate + w.volume + w.mix) - w.delta).abs().max() < 0.5
    return w


def stress_test(df24):
    """One-year stress on FY2024: +200bp cost of risk, -15% loans, -50bp NIM."""
    rows = {}
    for bank in BANKS:
        r = df24.loc[bank]
        nim = r["nii"] / r["avg_total_assets"]
        cor = -r["impairment"] / r["avg_loans"]
        cet1_capital = r["cet1"] / 100 * r["rwa"]
        loans_adv = r["loans"] * (1 + STRESS["loans"])
        assets_adv = r["avg_total_assets"] + r["loans"] * STRESS["loans"]
        nii_adv = (nim + STRESS["nim_bp"] / 10_000) * assets_adv
        imp_adv = (cor + STRESS["cor_bp"] / 10_000) * loans_adv
        pbt_adv = r["pbt"] + (nii_adv - r["nii"]) - (imp_adv + r["impairment"])
        np_adv = r["net_profit"] + (pbt_adv - r["pbt"]) * (1 - STRESS["tax"])
        rows[bank] = {
            "net_profit": r["net_profit"], "net_profit_stress": np_adv,
            "cet1_start": r["cet1"], "cet1_end": (cet1_capital + np_adv) / r["rwa"] * 100,
            "absorb_bp": r["pbt"] / r["loans"] * 10_000,
        }
    return pd.DataFrame(rows).T.loc[BANKS]


def justified_ptbv(rote):
    return (rote - G) / (COE - G)


# Cut-offs of the CAMELS-style score (1 = best, 5 = worst): the project's own, not a supervisory method
CAMELS_CUTS = {
    "cet1": ([11.0, 13.0, 15.0, 18.0], "higher"),
    "npe_ratio": ([2.0, 3.0, 5.0, 8.0], "lower"),
    "cost_to_income": ([30.0, 37.0, 45.0, 55.0], "lower"),
    "roe": ([5.0, 8.0, 12.0, 15.0], "higher"),
    "loan_to_deposit": ([60.0, 65.0, 75.0, 85.0], "lower"),
}


def camels_score(value, cuts, direction):
    t = sorted(cuts)
    if direction == "higher":
        return 1 + sum(value < x for x in reversed(t))
    return 1 + sum(value > x for x in t)


def camels_composite(df24):
    cols = {m: [camels_score(df24.loc[b, m], *CAMELS_CUTS[m]) for b in BANKS] for m in CAMELS_CUTS}
    return pd.Series(pd.DataFrame(cols, index=BANKS).mean(axis=1).round(1), index=BANKS)


def nii_2026_base_case():
    """NII 2026 = NIM x average assets. NIM = 2024 NIM + slope x (2026 rate - 2024 rate), the slope being
    the bank's 2022-2024 regression of NIM on the yearly average ECB deposit rate. Assets grow 2% a year
    from end-2024. Used only to check the typed base case values."""
    ecb = pd.read_csv(ECB_RATES_CSV).set_index("year")["dfr_avg"]
    out = {}
    for bank in BANKS:
        b = kpis[kpis.bank == bank].sort_values("year")
        slope = np.polyfit([ecb[y] for y in YEARS], b["nim"].tolist(), 1)[0]
        nim26 = b["nim"].iloc[-1] + slope * (2.00 - ecb[2024])
        assets26 = b["total_assets"].iloc[-1] * 1.02 * 1.01
        out[bank] = nim26 / 100 * assets26
    return pd.Series(out)


walk = rate_volume_walk(kpis)
stress = stress_test(k24)
walk_sector = walk.groupby("year")[["delta", "rate", "volume", "mix"]].sum()
nii_growth = (sector.loc[2024, "nii"] / sector.loc[2022, "nii"] - 1) * 100
np_2023_reported = (sector.loc[2023, "net_profit"] / sector.loc[2022, "net_profit"] - 1) * 100
np_2023_underlying = (sector.loc[2023, "net_profit"] / VERIFIED["sector_underlying_2022"] - 1) * 100
composite = camels_composite(k24)


def verify_results():
    """Stop the build if a recomputed result differs from the verified value."""
    V = VERIFIED
    for y in YEARS:
        assert round(sector.loc[y, "nii"]) == V["sector_nii"][y], f"sector NII {y}"
        assert round(sector.loc[y, "net_profit"]) == V["sector_net_profit"][y], f"sector profit {y}"
    assert round(nii_growth, 1) == V["sector_nii_growth_pct"]
    assert round(walk_sector.loc[2023, "delta"]) == V["walk_2023"]["delta"]
    assert round(walk_sector.loc[2023, "rate"]) == V["walk_2023"]["rate"]
    assert round(walk_sector.loc[2024, "delta"]) == V["walk_2024"]["delta"]
    assert round(walk_sector.loc[2024, "rate"]) == V["walk_2024"]["rate"]
    assert round(walk_sector.loc[2024, "volume"]) == V["walk_2024"]["volume"]
    eb = walk[(walk.bank == "Eurobank") & (walk.year == 2024)].iloc[0]
    assert round(eb["volume"]) == V["eurobank_volume_2024"]
    assert round(np_2023_reported, 1) == V["reported_change_2023_pct"]
    assert round(np_2023_underlying) == V["underlying_change_2023_pct"]
    for bank, (np_s, c0, c1, bp) in V["stress"].items():
        r = stress.loc[bank]
        assert round(r["net_profit_stress"]) == np_s, f"stress profit {bank}"
        assert round(r["cet1_start"], 1) == c0, f"stress CET1 start {bank}"
        assert round(r["cet1_end"], 1) == c1, f"stress CET1 end {bank}"
        assert round(r["absorb_bp"]) == bp, f"stress absorb {bank}"
    for y, v in V["dtc_total_bn"].items():
        assert round(sector.loc[y, "dtc"] / 1000, 1) == v, f"DTC total {y}"
    for bank, v in V["dtc_to_cet1_2024"].items():
        assert round(k24.loc[bank, "dtc_to_cet1"]) == v, f"DTC share {bank}"
    assert sum(V["nii_2026_base"].values()) - V["sector_nii_2026_base"] in (0, 1)   # rounding
    assert round((V["sector_nii_2026_base"] / sector.loc[2024, "nii"] - 1) * 100, 1) \
        == V["sector_nii_2026_vs_2024_pct"]
    base_check = nii_2026_base_case()
    for bank, v in V["nii_2026_base"].items():
        assert abs(base_check[bank] - v) <= 0.6, f"NII 2026 base {bank}"
    assert abs(base_check.sum() - V["sector_nii_2026_base"]) <= 0.6
    for bank, v in V["camels_2024"].items():
        assert abs(composite[bank] - v) < 1e-9, f"CAMELS composite {bank}"


verify_results()


# ---------------------------------------------------------------------------
# Claims in the text are checked against the data before they are written
# ---------------------------------------------------------------------------
def best_banks(df, metric, high=True):
    s = df[metric]
    target = s.max() if high else s.min()
    return [b for b in s.index if s[b] == target]


def is_best(metric, bank, high=True, df=k24):
    return bank in best_banks(df, metric, high)


def num(x, decimals=0):
    """Number as text, rounded half up like Excel (8,024.5 becomes 8,025)."""
    d = Decimal(str(float(x))).quantize(Decimal(1).scaleb(-decimals), rounding=ROUND_HALF_UP)
    return f"{d:,.{decimals}f}"


def eur(x, decimals=0):
    """EUR million as text, with a plain minus for negatives."""
    s = f"€{num(abs(x), decimals)}m"
    return f"-{s}" if x < 0 else s


def pct1(x):
    return f"{num(x, 1)}%"


eba_change = {b: VERIFIED["eba_2023"][b][1] - VERIFIED["eba_2023"][b][0] for b in BANKS}
nii_2026_change = {b: VERIFIED["nii_2026_base"][b] / k24.loc[b, "nii"] - 1 for b in BANKS}
ECB_FUNDING = VERIFIED["ecb_funding_bn"]

# ---------------------------------------------------------------------------
# Theme (dark navy, like the Streamlit app and the dashboard)
# ---------------------------------------------------------------------------
BG = RGBColor(10, 22, 40)
PANEL = RGBColor(30, 41, 59)
PANEL_DARK = RGBColor(15, 30, 50)
TEXT = RGBColor(241, 245, 249)
MUTED = RGBColor(166, 178, 196)
GOOD = RGBColor(74, 222, 128)
BAD = RGBColor(248, 113, 113)

# Brand colours for header bars. On the dark background the charts use lighter display colours
# (the NBG brand navy would almost disappear).
BRAND = {"Eurobank": (0, 103, 177), "Alpha Bank": (226, 35, 26), "Piraeus Bank": (247, 166, 0), "NBG": (0, 48, 135)}
DISPLAY = {"Eurobank": (46, 144, 232), "Alpha Bank": (240, 84, 72), "Piraeus Bank": (247, 166, 0),
           "NBG": (150, 160, 255)}


def rgb(t):
    return RGBColor(*t)


def hexc(t):
    return "#%02x%02x%02x" % tuple(t)


def mpl(t):
    return tuple(c / 255 for c in t)


MPL_THEME = {
    "figure.facecolor": "#0a1628",
    "axes.facecolor": "#0f1e30",
    "axes.edgecolor": "#334155",
    "axes.labelcolor": "#cbd5e1",
    "axes.titlecolor": "#f1f5f9",
    "axes.titlesize": 12,
    "axes.titleweight": "bold",
    "axes.unicode_minus": False,
    "xtick.color": "#cbd5e1",
    "ytick.color": "#cbd5e1",
    "text.color": "#e2e8f0",
    "grid.color": "#1e293b",
    "grid.alpha": 1.0,
    "lines.linewidth": 2.4,
    "font.family": "sans-serif",
    "font.sans-serif": ["Carlito", "Calibri", "DejaVu Sans"],
    "font.size": 12,
    "legend.frameon": False,
    "legend.fontsize": 11,
}
LABEL = "#f1f5f9"


plt.rcParams.update(MPL_THEME)


def new_fig(*args, **kwargs):
    return plt.subplots(*args, **kwargs)


def style_axes(ax, grid_axis="y"):
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    ax.grid(axis=grid_axis, alpha=0.6, linewidth=0.8)
    ax.set_axisbelow(True)
    ax.tick_params(length=0)


def euro_axis(ax):
    """Show euro amounts on the y axis with a thousands separator, like the labels on the bars."""
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:,.0f}"))


# ---------------------------------------------------------------------------
# Presentation helpers
# ---------------------------------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
TITLE_ONLY = prs.slide_layouts[5]
_slide_no = [0]


def _to_back(slide, shape):
    tree = slide.shapes._spTree
    tree.remove(shape._element)
    tree.insert(2, shape._element)


def add_rect(sl, x, y, w, h, color, name="Panel"):
    shape = sl.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    shape.name = name
    return shape


def _style_run(run, size, bold=False, color=TEXT, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = "Calibri"


def _frame(tf, anchor=MSO_ANCHOR.TOP, margins=(0.05, 0.03, 0.05, 0.03)):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]


def add_text(sl, text, x, y, w, h, size=12, bold=False, color=TEXT, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, italic=False, name="Text"):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    _frame(tb.text_frame, anchor)
    p = tb.text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _style_run(run, size, bold, color, italic)
    return tb


def _bullet(p, indent_in=0.2, char="•", color=None):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent_in))))
    pPr.set("indent", str(-int(Inches(indent_in))))
    if color is not None:
        clr = etree.SubElement(pPr, qn("a:buClr"))
        srgb = etree.SubElement(clr, qn("a:srgbClr"))
        srgb.set("val", "%02X%02X%02X" % tuple(color))
    font = etree.SubElement(pPr, qn("a:buFont"))
    font.set("typeface", "Arial")
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", char)


def add_paragraphs(sl, paras, x, y, w, h, size=12, color=MUTED, bullets=True, space_after=7,
                   bullet_color=None, name="Text"):
    """paras: list of str or (lead, rest). A lead is written in bold and white before the rest."""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    _frame(tb.text_frame)
    first = True
    for item in paras:
        p = tb.text_frame.paragraphs[0] if first else tb.text_frame.add_paragraph()
        first = False
        p.space_after = Pt(space_after)
        lead, rest = item if isinstance(item, tuple) else (None, item)
        if lead:
            r = p.add_run()
            r.text = lead + " "
            _style_run(r, size, True, TEXT)
        r = p.add_run()
        r.text = rest
        _style_run(r, size, False, color)
        if bullets:
            _bullet(p, color=bullet_color)
    return tb


def text_on(rgb_tuple):
    """Dark or white text, whichever reads better on the given fill."""
    r, g, b = rgb_tuple
    lum = 0.2126 * r + 0.7152 * g + 0.0722 * b
    return BG if lum > 150 else TEXT


def new_slide(title, subtitle=None, bar=PANEL, bar_text=TEXT, sub_color=MUTED):
    _slide_no[0] += 1
    sl = prs.slides.add_slide(TITLE_ONLY)
    sl.background.fill.solid()
    sl.background.fill.fore_color.rgb = BG
    bar_shape = add_rect(sl, 0, 0, 13.333, 1.1, bar, name="Header bar")
    _to_back(sl, bar_shape)
    t = sl.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.35), Inches(0.1), Inches(12.6), Inches(0.62)
    _frame(t.text_frame, MSO_ANCHOR.MIDDLE, (0.05, 0.0, 0.05, 0.0))
    t.text_frame.text = title
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    _style_run(p.runs[0], 26, True, bar_text)
    if subtitle:
        add_text(sl, subtitle, 0.35, 0.72, 12.6, 0.32, size=12, color=sub_color, name="Subtitle")
    footer(sl)
    return sl


def footer(sl):
    add_text(sl, f"Greek Banking Sector Analysis 2022–2024  |  Spyros Papastergiou  |  {NOT_ADVICE.rstrip('.')}  |  "
                 f"{_slide_no[0]} / {N_SLIDES}",
             0.35, 7.1, 12.6, 0.28, size=10, color=MUTED, align=PP_ALIGN.RIGHT, name="Footer")


def add_chart(sl, fig, x, y, w, h, alt, dpi=200, name="Chart"):
    """Place a matplotlib figure inside the box (x, y, w, h) without distorting it."""
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    buf.seek(0)
    with Image.open(buf) as im:
        pw, ph = im.size
    buf.seek(0)
    scale = min(w / pw, h / ph)
    iw, ih = pw * scale, ph * scale
    pic = sl.shapes.add_picture(buf, Inches(x + (w - iw) / 2), Inches(y), Inches(iw), Inches(ih))
    pic.name = name
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    return pic


NO_STYLE_NO_GRID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def add_table(sl, x, y, col_w, header, rows, size=12, row_h=0.4, aligns=None, colors=None, bolds=None,
              alt="", name="Table", header_size=None):
    """A real PowerPoint table. colors and bolds are optional lists of lists shaped like rows."""
    n_rows, n_cols = len(rows) + 1, len(header)
    gf = sl.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(sum(col_w)), Inches(row_h * n_rows))
    gf.name = name
    gf._element.nvGraphicFramePr.cNvPr.set("descr", alt)
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1")
    tblPr.set("bandRow", "0")
    tblPr.find(qn("a:tableStyleId")).text = NO_STYLE_NO_GRID
    aligns = aligns or ["l"] + ["c"] * (n_cols - 1)
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Inches(w)
    for i in range(n_rows):
        tbl.rows[i].height = Inches(row_h)

    def put(cell, text, sz, bold, color, align, fill):
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.08)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        r = p.add_run()
        r.text = str(text)
        _style_run(r, sz, bold, color)

    for j, h in enumerate(header):
        put(tbl.cell(0, j), h, header_size or size, True, TEXT, aligns[j], PANEL)
    for i, row in enumerate(rows, 1):
        fill = PANEL_DARK if i % 2 else RGBColor(20, 34, 54)
        for j, v in enumerate(row):
            color = colors[i - 1][j] if colors and colors[i - 1][j] is not None else TEXT
            bold = bolds[i - 1][j] if bolds else j == 0
            put(tbl.cell(i, j), v, size, bold, color, aligns[j], fill)
    return gf


# ===========================================================================
# SLIDE 1: COVER
# ===========================================================================
_slide_no[0] += 1
cover = prs.slides.add_slide(TITLE_ONLY)
cover.background.fill.solid()
cover.background.fill.fore_color.rgb = BG
for i, bank in enumerate(BANKS):
    add_rect(cover, 0.6 + i * 0.95, 1.35, 0.85, 0.12, rgb(DISPLAY[bank]), name=f"{bank} colour bar")
t = cover.shapes.title
t.left, t.top, t.width, t.height = Inches(0.55), Inches(1.7), Inches(12.2), Inches(1.0)
_frame(t.text_frame, MSO_ANCHOR.MIDDLE, (0.05, 0, 0.05, 0))
t.text_frame.text = "GREEK BANKING SECTOR ANALYSIS"
t.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
_style_run(t.text_frame.paragraphs[0].runs[0], 40, True, TEXT)
add_text(cover, "FY2022–2024", 0.6, 2.75, 6, 0.55, size=26, color=MUTED, name="Period")
add_text(cover, "Eurobank  ·  Alpha Bank  ·  Piraeus Bank  ·  NBG", 0.6, 3.4, 12, 0.45, size=18, color=TEXT,
         name="Banks")
add_text(cover, "Consolidated group figures: " + ", ".join(ENTITIES[b] for b in
                                                           ("Alpha Bank", "Eurobank", "NBG", "Piraeus Bank")),
         0.6, 3.9, 12.3, 0.6, size=12, color=MUTED, name="Entities")
add_text(cover, "Latest restated figures. Every number is traced to a PDF page.", 0.6, 4.5, 12, 0.35, size=12,
         color=MUTED, name="Basis")
add_text(cover, "Spyros Papastergiou", 0.6, 5.25, 6, 0.4, size=16, bold=True, color=TEXT, name="Author")
add_text(cover, f"{CONTACT_EMAIL}   |   {CONTACT_GITHUB}   |   {CONTACT_DASHBOARD}", 0.6, 5.7, 12.4, 0.35,
         size=12, color=MUTED, name="Contact")
add_text(cover, "For portfolio and educational purposes only. " + NOT_ADVICE, 0.6, 6.6, 12, 0.35, size=12,
         color=MUTED, italic=True, name="Disclaimer")

# ===========================================================================
# SLIDE 2: EXECUTIVE SUMMARY
# ===========================================================================
sl = new_slide("Executive Summary", "Four listed groups, FY2022–2024. €m unless stated.")

assert is_best("net_profit", "Eurobank")
assert min(stress["net_profit_stress"].items(), key=lambda kv: kv[1])[0] == "Alpha Bank"
assert stress["net_profit_stress"].drop("Alpha Bank").min() > 0
assert stress["absorb_bp"].idxmin() == "Alpha Bank"
other_bp = [int(round(stress.loc[b, "absorb_bp"])) for b in BANKS if b != "Alpha Bank"]
assert (min(other_bp), max(other_bp)) == (353, 420)

cards = [
    (f"Sector NII up {nii_growth:.1f}% in two years",
     f"Net interest income of the four banks rose from {eur(sector.loc[2022, 'nii'])} in 2022 to "
     f"{eur(sector.loc[2024, 'nii'])} in 2024. Most of it came in 2023, when wider margins added "
     f"{eur(walk_sector.loc[2023, 'rate'])} (rate effect). In 2024 NII grew by only "
     f"{eur(walk_sector.loc[2024, 'delta'])}: {eur(walk_sector.loc[2024, 'rate'])} from rate and "
     f"{eur(walk_sector.loc[2024, 'volume'])} from volume.", BANKS[0]),
    ("2022 profit includes one-offs",
     f"Sector net profit for 2022 was {eur(sector.loc[2022, 'net_profit'])} as reported and "
     f"{eur(VERIFIED['sector_underlying_2022'])} on an underlying basis. On underlying figures 2023 profit "
     f"rose about {np_2023_underlying:.0f}% instead of falling {abs(np_2023_reported):.1f}%.", BANKS[2]),
    ("Alpha is the weak point under stress",
     f"In a one-year stress test three banks stay profitable and Alpha loses "
     f"{eur(abs(stress.loc['Alpha Bank', 'net_profit_stress']))}. Alpha can absorb "
     f"{stress.loc['Alpha Bank', 'absorb_bp']:.0f}bp of extra cost of risk before a pre-tax loss, against "
     f"{min(other_bp)} to {max(other_bp)}bp for the other three.", BANKS[1]),
    ("Capital quality and funding",
     f"DTCs fell from €{sector.loc[2022, 'dtc'] / 1000:.1f}bn in 2022 to €{sector.loc[2024, 'dtc'] / 1000:.1f}bn "
     f"in 2024. At end-2024 they were {k24.loc['Eurobank', 'dtc_to_cet1']:.0f}% of CET1 capital at Eurobank and "
     f"{k24.loc['Piraeus Bank', 'dtc_to_cet1']:.0f}% at Piraeus. ECB funding of the four banks fell from "
     f"€{ECB_FUNDING[2022]}bn (end-2022) to €{ECB_FUNDING[2024]}bn (end-2024), all of it at Alpha.", BANKS[3]),
]
card_w, gap = 3.0, 0.2
for i, (head, body, bank) in enumerate(cards):
    cx = 0.35 + i * (card_w + gap)
    add_rect(sl, cx, 1.3, card_w, 2.65, PANEL_DARK, name="Card")
    add_rect(sl, cx, 1.3, card_w, 0.07, rgb(DISPLAY[bank]), name="Card accent")
    add_text(sl, head, cx + 0.1, 1.45, card_w - 0.2, 0.62, size=15, bold=True, color=TEXT, name="Card title")
    add_text(sl, body, cx + 0.1, 2.1, card_w - 0.2, 2.1, size=12.5, color=MUTED, name="Card text")

add_text(sl, "FY2024 key figures", 0.35, 4.15, 6, 0.3, size=14, bold=True, color=TEXT, name="Table title")
metrics = [("net_profit", True), ("roe", True), ("rote", True), ("nim", True), ("cost_to_income", False),
           ("cet1", True), ("npe_ratio", False), ("dtc_to_cet1", False)]
rows, colors = [], []
for bank in BANKS:
    r = k24.loc[bank]
    rows.append([bank, eur(r["net_profit"]), pct1(r["roe"]), pct1(r["rote"]), f"{r['nim']:.2f}%",
                 pct1(r["cost_to_income"]), pct1(r["cet1"]), pct1(r["npe_ratio"]), f"{r['dtc_to_cet1']:.0f}%"])
    c = [rgb(DISPLAY[bank])]
    for metric, high in metrics:
        c.append(GOOD if bank in best_banks(k24, metric, high) else None)
    colors.append(c)
add_table(sl, 0.35, 4.5, [1.9, 1.55, 1.15, 1.15, 1.15, 1.75, 1.2, 1.15, 1.6],
          ["Bank", "Net profit", "ROE", "RoTE", "NIM", "Cost-to-income", "CET1", "NPE ratio", "DTC / CET1"],
          rows, size=13, row_h=0.4, colors=colors,
          alt="Table of FY2024 net profit, ROE, RoTE, NIM, cost-to-income, CET1, NPE ratio and DTC share of CET1 "
              "for the four banks. Green marks the best of the four.", name="FY2024 key figures")
add_text(sl, "Green = best of the four (lowest for cost-to-income, NPE ratio and DTC / CET1). "
             "DTC = deferred tax credit, a tax claim on the Greek State that counts in CET1 capital.",
         0.35, 6.6, 12.6, 0.4, size=11, color=MUTED, name="Table note")

# ===========================================================================
# SLIDE 3: SECTOR LANDSCAPE
# ===========================================================================
sl = new_slide("Sector Landscape: NII and Profitability",
               f"Sector NII rose {nii_growth:.1f}% from 2022 to 2024. €m unless stated.")

fig, (ax, ax2) = new_fig(1, 2, figsize=(12.6, 4.5), gridspec_kw={"width_ratios": [1.15, 1]})
x = np.arange(len(BANKS))
bw = 0.26
for j, yr in enumerate(YEARS):
    vals = [kpis[(kpis.bank == b) & (kpis.year == yr)]["nii"].iloc[0] for b in BANKS]
    bars = ax.bar(x + (j - 1) * bw, vals, bw * 0.94, color=[mpl(DISPLAY[b]) for b in BANKS],
                  alpha=[0.45, 0.72, 1.0][j])
    for bar, v in zip(bars, vals):
        ax.text(bar.get_x() + bar.get_width() / 2, v + 40, f"{v:,.0f}", ha="center", va="bottom", fontsize=9,
                color=LABEL)
ax.set_title("Net interest income by bank (€m)", pad=10)
ax.set_xticks(x)
ax.set_xticklabels(BANKS, fontsize=11)
ax.set_ylim(0, 3150)
ax.legend(handles=[Patch(facecolor="#94a3b8", alpha=a, label=str(y)) for a, y in zip([0.45, 0.72, 1.0], YEARS)],
          loc="upper center", ncol=3, fontsize=10.5)
style_axes(ax)
euro_axis(ax)

label_y = {}
for bank in BANKS:
    b = kpis[kpis.bank == bank].sort_values("year")
    ax2.plot(b["year"], b["roe"], color=mpl(DISPLAY[bank]), marker="o", markersize=7)
    label_y[bank] = b["roe"].iloc[-1]
# spread the end labels so that close values do not overlap
order = sorted(BANKS, key=lambda k: label_y[k])
pos = {}
min_gap = 1.5
for i, bank in enumerate(order):
    pos[bank] = label_y[bank] if i == 0 else max(label_y[bank], pos[order[i - 1]] + min_gap)
for bank in BANKS:
    ax2.text(2024.08, pos[bank], f"{bank} {label_y[bank]:.1f}%", color=mpl(DISPLAY[bank]), fontsize=10.5,
             va="center", fontweight="bold")
ax2.axhline(COE, linestyle="--", color="#94a3b8", linewidth=1.2)
ax2.text(2021.9, COE + 0.3, f"Assumed cost of equity {COE:.0f}%", color="#cbd5e1", fontsize=10.5, va="bottom")
ax2.set_title("Return on average equity (%)", pad=10)
ax2.set_xticks(YEARS)
ax2.set_xlim(2021.8, 2025.1)
ax2.set_ylim(4, 24)
ax2.set_yticks([5, 10, 15, 20])
ax2.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: f"{v:.0f}%"))
style_axes(ax2)
fig.tight_layout(w_pad=2.5)
add_chart(sl, fig, 0.35, 1.25, 12.6, 4.6,
          f"Left: grouped bar chart of net interest income for each bank in 2022, 2023 and 2024. Sector total "
          f"{sector.loc[2022, 'nii']:,.0f}, {sector.loc[2023, 'nii']:,.0f} and {sector.loc[2024, 'nii']:,.0f} million "
          f"euro. Right: line chart of return on average equity per bank, 2022 to 2024, with a dashed line at the "
          f"assumed cost of equity of 10 percent.", name="NII and ROE charts")
add_paragraphs(sl, [
    f"Net profit of the four banks: {eur(sector.loc[2022, 'net_profit'])} (2022), "
    f"{eur(sector.loc[2023, 'net_profit'])} (2023), {eur(sector.loc[2024, 'net_profit'])} (2024). "
    f"2022 includes one-off items (underlying {eur(VERIFIED['sector_underlying_2022'])}).",
    "Eurobank 2024 includes Hellenic Bank, consolidated from 30 June 2024.",
], 0.35, 5.95, 12.6, 1.0, size=12, color=MUTED, space_after=4)

# ===========================================================================
# SLIDES 4 TO 7: ONE PAGE PER BANK
# ===========================================================================


def bank_slide(bank):
    r, p = k24.loc[bank], k23.loc[bank]
    br = BRAND[bank]
    sl = new_slide(bank, f"{ENTITIES[bank]}  |  FY2024 snapshot", bar=rgb(br), bar_text=text_on(br),
                   sub_color=text_on(br))
    cards = [
        ("ROE", pct1(r["roe"]), r["roe"] - p["roe"], 1, True),
        ("NIM", f"{r['nim']:.2f}%", r["nim"] - p["nim"], 2, True),
        ("CET1", pct1(r["cet1"]), r["cet1"] - p["cet1"], 1, True),
        ("Cost-to-income", pct1(r["cost_to_income"]), r["cost_to_income"] - p["cost_to_income"], 1, False),
        ("NPE ratio", pct1(r["npe_ratio"]), r["npe_ratio"] - p["npe_ratio"], 1, False),
        ("Net profit", eur(r["net_profit"]), (r["net_profit"] - p["net_profit"]) / abs(p["net_profit"]) * 100, None,
         True),
    ]
    cw, cg = 1.98, 0.14
    for i, (label, value, delta, dec, up_good) in enumerate(cards):
        cx = 0.35 + i * (cw + cg)
        add_rect(sl, cx, 1.3, cw, 1.15, PANEL_DARK, name="KPI card")
        add_rect(sl, cx, 1.3, cw, 0.06, rgb(DISPLAY[bank]), name="KPI accent")
        add_text(sl, label, cx + 0.05, 1.38, cw - 0.1, 0.28, size=11, color=MUTED, align=PP_ALIGN.CENTER,
                 name="KPI label")
        add_text(sl, value, cx + 0.05, 1.64, cw - 0.1, 0.45, size=20, bold=True, color=TEXT,
                 align=PP_ALIGN.CENTER, name="KPI value")
        if dec is None:
            txt = f"{delta:+.1f}% vs 2023"
        else:
            txt = f"{delta:+.{dec}f} pp vs 2023"
        good = (delta > 0) == up_good
        col = MUTED if abs(delta) < 1e-9 else (GOOD if good else BAD)
        add_text(sl, txt, cx + 0.05, 2.1, cw - 0.1, 0.28, size=11, color=col, align=PP_ALIGN.CENTER,
                 name="KPI change")
    return sl


def bank_chart(bank):
    b = kpis[kpis.bank == bank].sort_values("year")
    fig, ax1 = new_fig(figsize=(5.9, 3.7))
    col = mpl(DISPLAY[bank])
    bars = ax1.bar(b["year"], b["nii"], 0.5, color=col, alpha=0.8)
    for bar, v in zip(bars, b["nii"]):
        ax1.text(bar.get_x() + bar.get_width() / 2, v / 2, f"{v:,.0f}", ha="center", va="center", fontsize=11,
                 color="#0a1628", fontweight="bold")
    ax1.set_ylabel("NII (€m)")
    ax1.set_ylim(0, b["nii"].max() * 1.45)
    ax1.set_xticks(YEARS)
    style_axes(ax1)
    euro_axis(ax1)
    ax2 = ax1.twinx()
    ax2.plot(b["year"], b["roe"], color="#4ade80", marker="o", markersize=7)
    for yr, v in zip(b["year"], b["roe"]):
        ax2.annotate(f"{v:.1f}%", (yr, v), textcoords="offset points", xytext=(0, 10), ha="center", fontsize=11,
                     color="#4ade80", fontweight="bold",
                     bbox=dict(boxstyle="round,pad=0.18", fc="#0a1628", ec="none", alpha=0.9))
    ax2.set_ylim(0, max(b["roe"].max() * 1.35, 10))
    ax2.set_ylabel("ROE (%)", color="#4ade80")
    ax2.tick_params(axis="y", colors="#4ade80", length=0)
    for side in ("top",):
        ax2.spines[side].set_visible(False)
    ax2.spines["right"].set_visible(False)
    ax1.set_title("NII (bars) and ROE (line)", pad=10)
    fig.tight_layout()
    return fig


# Statements about each bank, checked against the data before they are written.
assert is_best("roe", "Eurobank") and is_best("net_profit", "Eurobank")
assert is_best("dtc_to_cet1", "Eurobank", high=False)
assert is_best("roe", "Alpha Bank", high=False) and is_best("rote", "Alpha Bank", high=False)
assert is_best("npe_ratio", "Alpha Bank", high=True)            # highest NPE ratio
assert is_best("cost_to_income", "Alpha Bank", high=True)       # highest cost-to-income
assert is_best("cost_of_risk", "Alpha Bank", high=True)         # highest cost of risk
assert is_best("cost_to_income", "Piraeus Bank", high=False)
assert is_best("cost_of_risk", "Piraeus Bank", high=False)
assert is_best("dtc_to_cet1", "Piraeus Bank", high=True)
assert is_best("cet1", "Piraeus Bank", high=False)
assert is_best("nim", "NBG") and is_best("cet1", "NBG")
assert is_best("loan_to_deposit", "NBG", high=False)
assert set(best_banks(k24, "npe_ratio", high=False)) == {"NBG", "Piraeus Bank"}
assert min(nii_2026_change, key=nii_2026_change.get) == "NBG"
assert min(eba_change, key=eba_change.get) == "Alpha Bank"
assert max(stress["absorb_bp"].items(), key=lambda kv: kv[1])[0] == "NBG"
vol24 = walk[walk.year == 2024].set_index("bank")["volume"]
rate24 = walk[walk.year == 2024].set_index("bank")["rate"]
assert vol24["Alpha Bank"] < 0
assert stress["absorb_bp"].idxmin() == "Alpha Bank" and stress["absorb_bp"].idxmax() == "NBG"


def stress_line(bank, rank_note=""):
    s = stress.loc[bank]
    cet1 = f"CET1 moves from {pct1(s['cet1_start'])} to {pct1(s['cet1_end'])}"
    tail = f"Extra cost of risk before a pre-tax loss: {s['absorb_bp']:.0f}bp{rank_note}."
    if s["net_profit_stress"] >= 0:
        return f"Net profit stays at {eur(s['net_profit_stress'])} and {cet1}. {tail}"
    return f"It is the only bank with a loss, of {eur(abs(s['net_profit_stress']))}, and {cet1}. {tail}"


def eba_line(bank):
    a, b = VERIFIED["eba_2023"][bank]
    return f"fully loaded CET1 fell from {a:.1f}% (Dec 2022) to {b:.1f}% (end-2025)"


BANK_POINTS = {
    "Eurobank": [
        ("Returns.", f"Highest ROE ({pct1(k24.loc['Eurobank', 'roe'])}) and largest net profit "
                     f"({eur(k24.loc['Eurobank', 'net_profit'])}) of the four banks in 2024."),
        ("Hellenic Bank.", f"Consolidated from 30 June 2024. Total assets rose {k24.loc['Eurobank', 'assets_yoy']:.1f}% "
                           f"to {eur(k24.loc['Eurobank', 'total_assets'])} and deposits rose "
                           f"{(k24.loc['Eurobank', 'deposits'] / k23.loc['Eurobank', 'deposits'] - 1) * 100:.1f}%. "
                           f"Its volume effect on NII was {eur(vol24['Eurobank'])}."),
        ("Capital quality.", f"DTCs were {k24.loc['Eurobank', 'dtc_to_cet1']:.0f}% of CET1 capital at end-2024, "
                             f"the lowest share of the four banks."),
        ("Stress test.", stress_line("Eurobank")),
    ],
    "Alpha Bank": [
        ("Returns.", f"Lowest ROE ({pct1(k24.loc['Alpha Bank', 'roe'])}) and RoTE "
                     f"({pct1(k24.loc['Alpha Bank', 'rote'])}) of the four banks in 2024."),
        ("NII.", f"Fell {abs(k24.loc['Alpha Bank', 'nii_yoy']):.1f}% in 2024 after rising "
                 f"{k23.loc['Alpha Bank', 'nii_yoy']:.1f}% in 2023. The volume effect took "
                 f"{eur(abs(vol24['Alpha Bank']))} off NII."),
        ("Asset quality and costs.", f"NPE ratio fell from {pct1(k22.loc['Alpha Bank', 'npe_ratio'])} (2022) to "
                                     f"{pct1(k24.loc['Alpha Bank', 'npe_ratio'])} (2024), still the highest of the "
                                     f"four. Cost-to-income ({pct1(k24.loc['Alpha Bank', 'cost_to_income'])}) and "
                                     f"cost of risk ({k24.loc['Alpha Bank', 'cost_of_risk']:.2f}%) are the highest too."),
        ("Funding.", f"The only bank still funded by the ECB at end-2024 (€{ECB_FUNDING[2024]}bn)."),
        ("Stress test.", stress_line("Alpha Bank", ", the least of the four banks")),
        ("EBA 2023.", f"In the official adverse scenario {eba_line('Alpha Bank')}, the largest fall of the four."),
    ],
    "Piraeus Bank": [
        ("Efficiency.", f"Lowest cost-to-income ({pct1(k24.loc['Piraeus Bank', 'cost_to_income'])}) and lowest cost "
                        f"of risk ({k24.loc['Piraeus Bank', 'cost_of_risk']:.2f}%) of the four banks in 2024."),
        ("Profit.", f"Net profit rose {k24.loc['Piraeus Bank', 'profit_yoy']:.1f}% to "
                    f"{eur(k24.loc['Piraeus Bank', 'net_profit'])}, while NII grew only "
                    f"{k24.loc['Piraeus Bank', 'nii_yoy']:.1f}%. Loan impairments fell to "
                    f"{eur(abs(k24.loc['Piraeus Bank', 'impairment']))} from "
                    f"{eur(abs(k23.loc['Piraeus Bank', 'impairment']))}."),
        ("Capital quality.", f"DTCs were {k24.loc['Piraeus Bank', 'dtc_to_cet1']:.0f}% of CET1 capital at end-2024, "
                             f"the highest share of the four banks. CET1 ratio of "
                             f"{pct1(k24.loc['Piraeus Bank', 'cet1'])} is the lowest of the four."),
        ("Stress test.", stress_line("Piraeus Bank")),
    ],
    "NBG": [
        ("Margin and capital.", f"Highest NIM ({k24.loc['NBG', 'nim']:.2f}%) and highest CET1 ratio "
                                f"({pct1(k24.loc['NBG', 'cet1'])}) of the four banks in 2024."),
        ("Asset quality and liquidity.", f"NPE ratio of {pct1(k24.loc['NBG', 'npe_ratio'])}, level with Piraeus. "
                                         f"Loan-to-deposit ratio of {pct1(k24.loc['NBG', 'loan_to_deposit'])}, the "
                                         f"lowest of the four."),
        ("Rough guide to 2026.", f"In the base case NII falls {abs(nii_2026_change['NBG']) * 100:.1f}% from 2024 "
                                 f"to {eur(VERIFIED['nii_2026_base']['NBG'])}, the largest fall of the four. "
                                 f"This is not a forecast."),
        ("Stress test.", stress_line("NBG", ", the most of the four banks")),
    ],
}

for n, bank in enumerate(BANKS):
    sl = bank_slide(bank)
    add_text(sl, "What the numbers show", 0.35, 2.7, 6.6, 0.35, size=15, bold=True, color=TEXT, name="Points title")
    add_rect(sl, 0.35, 3.07, 6.6, 0.03, rgb(DISPLAY[bank]), name="Points rule")
    add_paragraphs(sl, BANK_POINTS[bank], 0.35, 3.2, 6.6, 3.85, size=14 if len(BANK_POINTS[bank]) <= 4 else 13,
                   color=MUTED, space_after=8,
                   bullet_color=DISPLAY[bank], name="Points")
    b = kpis[kpis.bank == bank].sort_values("year")
    add_chart(sl, bank_chart(bank), 7.2, 2.65, 5.75, 4.3,
              f"{bank}: bars show net interest income of {b['nii'].iloc[0]:,.0f}, {b['nii'].iloc[1]:,.0f} and "
              f"{b['nii'].iloc[2]:,.0f} million euro in 2022, 2023 and 2024. The line shows return on equity of "
              f"{b['roe'].iloc[0]:.1f}, {b['roe'].iloc[1]:.1f} and {b['roe'].iloc[2]:.1f} percent.",
              name=f"{bank} chart")

# ===========================================================================
# SLIDE 8: PEER RANKINGS
# ===========================================================================
sl = new_slide("Peer Rankings: FY2024",
               "Green outline = best of the four. Lower is better for cost-to-income and the NPE ratio.")
ranking = [("roe", "ROE (%)", True, "%"), ("nim", "NIM (%)", True, "%"), ("cet1", "CET1 ratio (%)", True, "%"),
           ("cost_to_income", "Cost-to-income (%)", False, "%"), ("npe_ratio", "NPE ratio (%)", False, "%"),
           ("net_profit", "Net profit (€m)", True, "")]
fig, axes = new_fig(2, 3, figsize=(12.6, 5.4))
for ax, (key, label, high, sfx) in zip(axes.flat, ranking):
    vals = k24[key]
    order = vals.sort_values(ascending=not high).index.tolist()
    best = best_banks(k24, key, high)
    ypos = np.arange(len(order))
    bars = ax.barh(ypos, [vals[b] for b in order], color=[mpl(DISPLAY[b]) for b in order], height=0.62)
    for bar, b in zip(bars, order):
        if b in best:
            bar.set_edgecolor("#4ade80")
            bar.set_linewidth(2.2)
        v = vals[b]
        txt = f"{v:.2f}%" if key == "nim" else (f"{v:.1f}%" if sfx else f"{v:,.0f}")
        ax.text(v + vals.max() * 0.02, bar.get_y() + bar.get_height() / 2, txt, va="center", fontsize=10.5,
                color=LABEL)
    ax.set_yticks(ypos)
    ax.set_yticklabels(order, fontsize=10.5)
    ax.invert_yaxis()
    ax.set_xlim(0, vals.max() * 1.28)
    ax.set_xticks([])
    ax.set_title(label, fontsize=12, pad=6)
    ax.grid(False)
    for side in ("top", "right", "bottom"):
        ax.spines[side].set_visible(False)
    ax.tick_params(length=0)
fig.tight_layout(h_pad=2.2, w_pad=1.6)
leaders = {label: " and ".join(best_banks(k24, key, high)) for key, label, high, _ in ranking}
add_chart(sl, fig, 0.35, 1.25, 12.6, 5.7,
          "Six bar charts of FY2024 values per bank, best on top: "
          + "; ".join(f"{label.split(' (')[0]}: {leaders[label]}" for _, label, _, _ in ranking) + ".",
          name="Peer rankings")

# ===========================================================================
# SLIDE 9: NII RATE VS VOLUME AND ROUGH GUIDE TO 2026
# ===========================================================================
sl = new_slide("NII: Rate vs Volume, and a Rough Guide to 2026",
               "Sector NII walk 2022 to 2024 and a base case for 2026. €m. A rough guide, not a forecast.")

fig, (ax, ax2) = new_fig(1, 2, figsize=(12.6, 4.2), gridspec_kw={"width_ratios": [1.35, 1]})
W23, W24 = walk_sector.loc[2023], walk_sector.loc[2024]
steps = [("NII 2022", sector.loc[2022, "nii"], "total"),
         ("Rate", W23["rate"], "rate"), ("Volume", W23["volume"], "volume"), ("Mix", W23["mix"], "mix"),
         ("NII 2023", sector.loc[2023, "nii"], "total"),
         ("Rate", W24["rate"], "rate"), ("Volume", W24["volume"], "volume"), ("Mix", W24["mix"], "mix"),
         ("NII 2024", sector.loc[2024, "nii"], "total")]
STEP_COLORS = {"total": "#cbd5e1", "rate": "#60a5fa", "volume": "#2dd4bf", "mix": "#a78bfa"}
level = 0.0
for i, (lab, v, kind) in enumerate(steps):
    if kind == "total":
        ax.bar(i, v, 0.62, color=STEP_COLORS[kind])
        ax.text(i, v + 130, f"{v:,.0f}", ha="center", va="bottom", fontsize=10.5, color=LABEL, fontweight="bold")
        level = v
    else:
        bottom = level if v >= 0 else level + v
        ax.bar(i, abs(v), 0.62, bottom=bottom, color=STEP_COLORS[kind])
        top = max(level, level + v)
        ax.text(i, top + 130, f"{v:+,.0f}", ha="center", va="bottom", fontsize=10.5, color=LABEL)
        level += v
ax.set_xticks(range(len(steps)))
ax.set_xticklabels([s[0] for s in steps], fontsize=10.5)
ax.set_ylim(0, 10200)
ax.set_title("Sector NII walk (€m)", pad=10)
style_axes(ax)
euro_axis(ax)

x = np.arange(len(BANKS))
bw = 0.36
v24 = [k24.loc[b, "nii"] for b in BANKS]
v26 = [VERIFIED["nii_2026_base"][b] for b in BANKS]
b1 = ax2.bar(x - bw / 2, v24, bw * 0.95, color=[mpl(DISPLAY[b]) for b in BANKS], alpha=1.0)
b2 = ax2.bar(x + bw / 2, v26, bw * 0.95, color=[mpl(DISPLAY[b]) for b in BANKS], alpha=0.45, hatch="//",
             edgecolor="#0a1628", linewidth=0)
for bar, v in zip(list(b1) + list(b2), v24 + v26):
    ax2.text(bar.get_x() + bar.get_width() / 2, v + 40, f"{v:,.0f}", ha="center", va="bottom", fontsize=9.5,
             color=LABEL)
ax2.set_xticks(x)
ax2.set_xticklabels([f"{b}\n{nii_2026_change[b] * 100:+.1f}%" for b in BANKS], fontsize=10.5)
ax2.set_ylim(0, 3350)
ax2.set_title("NII 2024 and 2026 base case (€m)", pad=10)
ax2.legend(handles=[Patch(facecolor="#94a3b8", label="2024 actual"),
                    Patch(facecolor="#94a3b8", alpha=0.45, hatch="//", label="2026 base case")],
           loc="upper center", ncol=2, fontsize=10.5)
style_axes(ax2)
euro_axis(ax2)
fig.tight_layout(w_pad=2.5)
add_chart(sl, fig, 0.35, 1.25, 12.6, 4.1,
          f"Left: waterfall of sector net interest income from {sector.loc[2022, 'nii']:,.0f} in 2022 to "
          f"{sector.loc[2023, 'nii']:,.0f} in 2023 and {sector.loc[2024, 'nii']:,.0f} in 2024, split into rate, volume "
          f"and mix effects. Right: net interest income per bank in 2024 and in the 2026 base case, with the "
          f"percentage change.", name="NII walk and 2026 base case")
add_paragraphs(sl, [
    ("2022 to 2023.", f"Wider margins (rate effect {eur(W23['rate'])}) drove NII up by {eur(W23['delta'])}. "
                      f"Volume and mix took a little back."),
    ("2023 to 2024.", f"NII rose by {eur(W24['delta'])}: rate {eur(W24['rate'])} and volume {eur(W24['volume'])}. "
                      f"Eurobank's volume effect was {eur(vol24['Eurobank'])}, after it consolidated Hellenic Bank "
                      f"from 30 June 2024. Rounded effects can differ from the total by €1m."),
    ("2026 base case.", f"Sector NII of {eur(VERIFIED['sector_nii_2026_base'])}, "
                        f"{abs(VERIFIED['sector_nii_2026_vs_2024_pct']):.1f}% below 2024. The percentage under each "
                        f"bank is its change from 2024. NIM moves with the ECB yearly average deposit rate (gradual "
                        f"cuts) and total assets grow 2% a year. The rate sensitivity of each bank rests on three "
                        f"yearly points, so this is directional only. Banks may not add up to the sector total "
                        f"because of rounding."),
], 0.35, 5.4, 12.6, 1.65, size=11.5, color=MUTED, space_after=4)

# ===========================================================================
# SLIDE 10: ONE-YEAR STRESS TEST
# ===========================================================================
sl = new_slide("One-Year Stress Test",
               "FY2024 base. Cost of risk up 200bp, loans down 15%, NIM down 50bp. Directional only.")
fig, (ax, ax2) = new_fig(1, 2, figsize=(12.6, 3.5))
x = np.arange(len(BANKS))
bw = 0.36
base_np = [stress.loc[b, "net_profit"] for b in BANKS]
adv_np = [stress.loc[b, "net_profit_stress"] for b in BANKS]
c1 = ax.bar(x - bw / 2, base_np, bw * 0.95, color=[mpl(DISPLAY[b]) for b in BANKS], alpha=1.0)
c2 = ax.bar(x + bw / 2, adv_np, bw * 0.95, color=[mpl(DISPLAY[b]) for b in BANKS], alpha=0.45, hatch="//",
            edgecolor="#0a1628", linewidth=0)
for bar, v in zip(list(c1) + list(c2), base_np + adv_np):
    ax.text(bar.get_x() + bar.get_width() / 2, v + (40 if v >= 0 else -40), f"{v:,.0f}", ha="center",
            va="bottom" if v >= 0 else "top", fontsize=10, color=LABEL)
ax.axhline(0, color="#64748b", linewidth=1)
ax.set_xticks(x)
ax.set_xticklabels(BANKS, fontsize=10.5)
ax.set_ylim(-420, 2050)
ax.set_title("Net profit: FY2024 and stressed (€m)", pad=10)
ax.legend(handles=[Patch(facecolor="#94a3b8", label="FY2024"),
                   Patch(facecolor="#94a3b8", alpha=0.45, hatch="//", label="Stressed")], loc="upper right",
          ncol=2)
style_axes(ax)
euro_axis(ax)
c0 = [stress.loc[b, "cet1_start"] for b in BANKS]
c3 = [stress.loc[b, "cet1_end"] for b in BANKS]
d1 = ax2.bar(x - bw / 2, c0, bw * 0.95, color=[mpl(DISPLAY[b]) for b in BANKS], alpha=1.0)
d2 = ax2.bar(x + bw / 2, c3, bw * 0.95, color=[mpl(DISPLAY[b]) for b in BANKS], alpha=0.45, hatch="//",
             edgecolor="#0a1628", linewidth=0)
for bar, v in zip(list(d1) + list(d2), c0 + c3):
    ax2.text(bar.get_x() + bar.get_width() / 2, v + 0.3, f"{v:.1f}%", ha="center", va="bottom", fontsize=10,
             color=LABEL)
ax2.set_xticks(x)
ax2.set_xticklabels(BANKS, fontsize=10.5)
ax2.set_ylim(0, 27)
ax2.set_title("CET1 ratio: end-2024 and end of stress year (%)", pad=10)
ax2.legend(handles=[Patch(facecolor="#94a3b8", label="End-2024"),
                    Patch(facecolor="#94a3b8", alpha=0.45, hatch="//", label="End of stress year")],
           loc="upper right", ncol=2)
style_axes(ax2)
fig.tight_layout(w_pad=2.5)
add_chart(sl, fig, 0.35, 1.2, 12.6, 3.3,
          "Left: net profit per bank in FY2024 and in the stress case. Alpha Bank turns to a loss of "
          f"{abs(stress.loc['Alpha Bank', 'net_profit_stress']):,.0f} million euro, the other three stay "
          "profitable. Right: CET1 ratio per bank at end-2024 and at the end of the stress year.",
          name="Stress test charts")

rows = []
for bank in BANKS:
    a, b = VERIFIED["eba_2023"][bank]
    rows.append([bank, f"{stress.loc[bank, 'absorb_bp']:.0f}bp", f"{a:.1f}%", f"{b:.1f}%", f"{b - a:+.1f} pp"])
colors = [[rgb(DISPLAY[b]), None, None, None, None] for b in BANKS]
add_table(sl, 0.35, 4.6, [2.2, 3.6, 2.3, 2.6, 1.9],
          ["Bank", "Extra cost of risk before a pre-tax loss", "EBA 2023: Dec 2022", "EBA 2023: adverse end-2025",
           "Change"], rows, size=11.5, row_h=0.34, colors=colors,
          alt="Table per bank: extra cost of risk the bank can absorb before a pre-tax loss in the one-year test, "
              "and the official EBA 2023 stress test result for fully loaded CET1.", name="Stress and EBA table")
add_text(sl, "Costs flat, 29% tax on the change in profit, no dividends, RWAs unchanged. Stressed CET1 = (reported CET1 "
             "capital + stressed profit) / RWAs, so a bank that stays profitable ends the year with more CET1 because "
             "no dividend is paid. The EBA 2023 test is a three-year scenario on the 31.12.2022 balance sheet and "
             "cannot be compared with the simplified test.", 0.35, 6.35, 12.6, 0.7, size=10.5, color=MUTED,
         name="Stress note")

# ===========================================================================
# SLIDE 11: VALUATION FRAMEWORK
# ===========================================================================
sl = new_slide("Valuation Framework: Justified P/TBV",
               f"Justified P/TBV = (RoTE - g) / (CoE - g), with CoE {COE:.0f}% and g {G:.0f}% as assumptions. "
               f"{NOT_ADVICE}")
jp = {b: justified_ptbv(k24.loc[b, "rote"]) for b in BANKS}
above = [b for b in BANKS if k24.loc[b, "rote"] > COE]
below = [b for b in BANKS if k24.loc[b, "rote"] <= COE]
assert above == ["Eurobank", "Piraeus Bank", "NBG"] and below == ["Alpha Bank"]
assert max(jp, key=jp.get) == "Eurobank" and jp["Alpha Bank"] < 1.0 <= min(jp[b] for b in above)

rows, colors = [], []
for bank in BANKS:
    r = k24.loc[bank]
    reading = "Earns above the assumed cost of equity" if bank in above else \
        "Earns slightly below the assumed cost of equity"
    rows.append([bank, eur(r["net_profit"]), eur(r["avg_tangible_equity"]), pct1(r["rote"]), f"{jp[bank]:.2f}x",
                 reading])
    colors.append([rgb(DISPLAY[bank]), None, None, None, TEXT, GOOD if bank in above else BAD])
add_table(sl, 0.35, 1.3, [1.9, 1.8, 2.1, 1.4, 1.7, 3.7],
          ["Bank", "Net profit 2024", "Avg tangible equity", "RoTE 2024", "Justified P/TBV", "Reading"],
          rows, size=13, row_h=0.42, colors=colors, aligns=["l", "c", "c", "c", "c", "l"],
          alt="Table per bank: net profit, average tangible equity, RoTE and justified price to tangible book "
              "value for FY2024.", name="Valuation table")

fig, ax = new_fig(figsize=(5.2, 3.3))
order = sorted(BANKS, key=lambda b: -jp[b])
bars = ax.bar(range(len(order)), [jp[b] for b in order], 0.55, color=[mpl(DISPLAY[b]) for b in order])
for bar, b in zip(bars, order):
    ax.text(bar.get_x() + bar.get_width() / 2, jp[b] + 0.05, f"{jp[b]:.2f}x", ha="center", va="bottom",
            fontsize=11.5, color=LABEL, fontweight="bold")
ax.axhline(1.0, linestyle="--", color="#94a3b8", linewidth=1.3)
ax.legend(handles=[Line2D([0], [0], linestyle="--", color="#94a3b8", label="1.0x: earns exactly the cost of equity")],
          loc="upper right", fontsize=10)
ax.set_xticks(range(len(order)))
ax.set_xticklabels(order, fontsize=10.5)
ax.set_ylim(0, max(jp.values()) * 1.3)
ax.set_title("Justified P/TBV", pad=8)
style_axes(ax)
fig.tight_layout()
add_chart(sl, fig, 7.6, 3.5, 5.35, 3.5,
          "Bar chart of justified price to tangible book value: "
          + ", ".join(f"{b} {jp[b]:.2f}x" for b in order)
          + ". A dashed line marks 1.0x, where a bank earns exactly its cost of equity.", name="Justified P/TBV chart")

add_paragraphs(sl, [
    ("What it measures.", "The multiple of tangible book value that a bank's return would justify if it lasted. "
                          "A bank that earns exactly its cost of equity is worth 1.0x."),
    ("What it does not do.", "No market prices are used, so the slide does not say whether any bank is cheap or "
                             "expensive. It is a way to read returns, not a price target."),
    ("Assumptions.", f"Cost of equity {COE:.0f}% and long-run growth {G:.0f}% are stated assumptions. Change them and "
                     f"the multiples change. RoTE = net profit / average (equity less AT1 less goodwill and "
                     f"intangibles)."),
], 0.35, 3.6, 7.0, 3.4, size=13, color=MUTED, space_after=9, bullets=True)

# ===========================================================================
# SLIDE 12: METHODOLOGY AND DATA QUALITY
# ===========================================================================
sl = new_slide("Methodology and Data Quality",
               "Annual reports only. Report file and PDF page recorded for every figure.")
cols = [
    ("Data", [
        "12 annual reports: FY2022, FY2023 and FY2024 of the four groups",
        "2024 and 2023 from the FY2024 reports (2023 restated), 2022 from the FY2023 reports",
        "Statements read with Python (pdfplumber). CET1, NPE, DTC and RWAs as reported at each year end",
    ]),
    ("Checks", [
        "Automated tests recompute ROE, ROA, NIM, cost of risk, cost-to-income and loan-to-deposit from the "
        "statement lines",
        "Assets equal liabilities plus equity for every bank and year",
        "Every statement line and ratio has a report page",
        "DTC share of CET1 is checked against the figure disclosed by Alpha and Eurobank",
    ]),
    ("Methods", [
        "DuPont breakdown of ROE",
        "CAMELS-style scorecard, own cut-offs, not a supervisory rating",
        "NII rate, volume and mix walk",
        "One-year stress test and the official EBA 2023 results",
        "Justified P/TBV with stated assumptions",
    ]),
    ("Limits", [
        "Yearly data only, so no view inside the year",
        "CET1 ratios are as reported: Eurobank's 16.8% is before its FY2024 dividend accrual, the other three include FY2024 profit. NPE ratios follow each bank's own definition",
        "Operating income and expenses follow each bank's own presentation",
        "Stress test and NII scenario are simple and static: directional only",
    ]),
]
col_w, gap = 3.0, 0.2
for i, (head, items) in enumerate(cols):
    cx = 0.35 + i * (col_w + gap)
    add_rect(sl, cx, 1.3, col_w, 0.42, PANEL, name="Column header")
    add_text(sl, head, cx + 0.1, 1.33, col_w - 0.2, 0.36, size=14, bold=True, color=TEXT,
             anchor=MSO_ANCHOR.MIDDLE, name="Column title")
    add_paragraphs(sl, items, cx, 1.85, col_w, 2.9, size=12.5, color=MUTED, space_after=7,
                   bullet_color=DISPLAY[BANKS[i]], name="Column text")

add_rect(sl, 0.35, 4.82, 12.6, 0.65, PANEL_DARK, name="Trace box")
add_rect(sl, 0.35, 4.82, 0.07, 0.65, rgb(DISPLAY["Eurobank"]), name="Trace accent")
add_paragraphs(sl, [("To check a number:", "find it in sources.csv in the repository. It lists the report file and the "
                                            "PDF page of every figure used here.")],
               0.55, 4.93, 12.2, 0.45, size=14, color=MUTED, bullets=False, space_after=0, name="Trace text")
add_rect(sl, 0.35, 5.62, 12.6, 0.03, rgb(DISPLAY["Eurobank"]), name="Rule")
add_text(sl, "Entities (consolidated group figures)", 0.35, 5.72, 6, 0.3, size=12, bold=True, color=TEXT,
         name="Entities title")
add_text(sl, ", ".join(ENTITIES[b] for b in ("Alpha Bank", "Eurobank", "NBG", "Piraeus Bank")),
         0.35, 6.0, 12.6, 0.4, size=12, color=MUTED, name="Entities")
add_text(sl, f"Spyros Papastergiou   |   {CONTACT_EMAIL}   |   {CONTACT_GITHUB}   |   {CONTACT_DASHBOARD}",
         0.35, 6.5, 12.6, 0.3, size=12, color=TEXT, name="Contact")
add_text(sl, "For portfolio and educational purposes only. " + NOT_ADVICE, 0.35, 6.78, 12.6, 0.3, size=11,
         color=MUTED, italic=True, name="Disclaimer")

assert len(prs.slides) == N_SLIDES, len(prs.slides)

# ---------------------------------------------------------------------------
# Document properties and save
# ---------------------------------------------------------------------------
cp = prs.core_properties
cp.title = "Greek Banking Sector Analysis 2022–2024"
cp.author = "Spyros Papastergiou"
cp.last_modified_by = "Spyros Papastergiou"
cp.subject = "Eurobank, Alpha Bank, Piraeus Bank and NBG, FY2022–2024. Not investment advice."
cp.keywords = "Greek banks, net interest income, stress test, CET1, DTC, justified P/TBV"
cp.comments = f"Built by _build_pptx.py from kpis_final.csv. Contact: {CONTACT_EMAIL}"
cp.created = cp.modified = datetime.now()

# The default template uses an en dash as a bullet character in its master text styles. Use a round bullet.
for part in [prs.slide_master] + list(prs.slide_layouts):
    for bu in part._element.iter(qn("a:buChar")):
        bu.set("char", "•")

prs.save(OUT)
print(f"Written: {OUT}  ({len(prs.slides)} slides)")
